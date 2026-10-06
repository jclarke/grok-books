"""Vendor aliases. Ledger rows stay as imported. Temporary database and key only."""

from __future__ import annotations

import io

import pytest
from openpyxl import load_workbook

from hpbooks.analytics import global_search, vendor_report
from hpbooks.cli import main
from hpbooks.db import HpbooksError, connect, init_db
from fake_accounts import BANK
from hpbooks.reports import query_transactions
from hpbooks.vendors import canonical_name, merge_vendors, rename_canonical, suggest_duplicates, unmerge_vendor

KEY = "ab" * 32
INTUIT_A = "ACH HOLD INTUIT PAYROLL"
INTUIT_B = "INTUIT PAYROLL LONG DESCRIPTOR"
OTHER = "ZZZ OTHER"


@pytest.fixture()
def env(tmp_path, monkeypatch):
    monkeypatch.setenv("HPBOOKS_KEY", KEY)
    monkeypatch.setenv("HPBOOKS_DB", str(tmp_path / "books.db"))
    monkeypatch.delenv("HPBOOKS_KEY_FILE", raising=False)
    init_db()
    return tmp_path


def _insert(conn, txn_id, day, cents, merchant, *, tag="general", category="Software & Licenses"):
    now = "2026-06-01T00:00:00+00:00"
    conn.execute(
        """
        INSERT INTO transactions (
          id, account_id, date, amount_cents, direction, currency,
          name, merchant_name, description, pending, status,
          first_seen_at, last_seen_at, updated_at, source
        ) VALUES (?, ?, ?, ?, 'out', 'USD', ?, ?, '', 0, 'active', ?, ?, ?, 'test')
        """,
        (txn_id, BANK, day, cents, merchant, merchant, now, now, now),
    )
    conn.execute(
        """
        INSERT INTO classifications (txn_id, business_tag, category, source, confidence, updated_at)
        VALUES (?, ?, ?, 'manual', 1, ?)
        """,
        (txn_id, tag, category, now),
    )


def _book(conn):
    _insert(conn, "a", "2026-06-02", -1000, INTUIT_A)
    _insert(conn, "b", "2026-06-03", -2000, INTUIT_B)
    _insert(conn, "c", "2026-06-04", -500, OTHER)


def _names(conn):
    rows, _table = vendor_report(conn, "2026-06-01", "2026-06-30", "all", limit=20)
    return rows


def test_merge_combines_totals_and_unmerge_splits_them(env):
    with connect() as conn:
        _book(conn)
        before = _names(conn)
        assert sum(row["spend_cents"] for row in before) == 3500
        assert {row["merchant"] for row in before} == {INTUIT_A, INTUIT_B, OTHER}

        merged = merge_vendors(conn, [INTUIT_A, INTUIT_B], "Intuit Payroll", actor="test")
        assert merged["changed"] == 2
        after = _names(conn)
        assert sum(row["spend_cents"] for row in after) == 3500
        intuit = next(row for row in after if row["merchant"] == "Intuit Payroll")
        assert intuit["spend_cents"] == 3000
        assert intuit["count"] == 2
        assert intuit["alias_count"] == 2
        assert intuit["spellings"] == [INTUIT_A, INTUIT_B]
        assert OTHER in {row["merchant"] for row in after}
        assert INTUIT_A not in {row["merchant"] for row in after}

        stored = {
            row["id"]: row["merchant_name"]
            for row in conn.execute("SELECT id, merchant_name FROM transactions")
        }
        assert stored == {"a": INTUIT_A, "b": INTUIT_B, "c": OTHER}

        again = merge_vendors(conn, [INTUIT_A, "Intuit Payroll"], "Intuit Payroll", actor="test")
        assert again["changed"] == 0
        audits = conn.execute(
            "SELECT action, actor, note FROM audit_log WHERE action = 'vendor_merge'"
        ).fetchall()
        assert len(audits) == 1
        assert audits[0]["actor"] == "test"
        assert audits[0]["note"] == "Intuit Payroll"

        rows, _total = query_transactions(conn, vendor="Intuit Payroll", date_from="2026-06-01", date_to="2026-06-30")
        assert {row["id"] for row in rows} == {"a", "b"}
        rows, _total = query_transactions(conn, vendor=INTUIT_A, date_from="2026-06-01", date_to="2026-06-30")
        assert {row["id"] for row in rows} == {"a", "b"}

        found = global_search(conn, "Payroll Processor")
        assert found["vendors"] == []
        rename_canonical(conn, "Intuit Payroll", "Payroll Processor", actor="test")
        found = global_search(conn, "Payroll Processor")
        assert found["vendors"][0]["name"] == "Payroll Processor"
        assert found["vendors"][0]["spend_cents"] == 3000
        renamed = next(row for row in _names(conn) if row["merchant"] == "Payroll Processor")
        assert renamed["spend_cents"] == 3000
        assert INTUIT_A in renamed["spellings"]

        unmerge_vendor(conn, INTUIT_A, actor="test")
        split = _names(conn)
        alone = next(row for row in split if row["merchant"] == INTUIT_A)
        still = next(row for row in split if row["merchant"] == "Payroll Processor")
        assert alone["spend_cents"] == 1000
        assert still["spend_cents"] == 2000
        assert sum(row["spend_cents"] for row in split) == 3500
        actions = [row["action"] for row in conn.execute("SELECT action FROM audit_log ORDER BY id")]
        assert actions.count("vendor_merge") == 1
        assert actions.count("vendor_unmerge") == 1
        assert actions.count("vendor_rename") == 1
        with pytest.raises(HpbooksError, match="no alias"):
            unmerge_vendor(conn, INTUIT_A, actor="test")
        with pytest.raises(HpbooksError, match="rename the canonical"):
            rename_canonical(conn, INTUIT_B, "Somewhere Else", actor="test")


def test_suggestions_and_idempotent_group_merge(env):
    with connect() as conn:
        _book(conn)
        from hpbooks.vendors import load_aliases

        spending = {INTUIT_A: 1000, INTUIT_B: 2000, OTHER: 500}
        ideas = suggest_duplicates(spending, load_aliases(conn))
        pair = next(item for item in ideas if set(item["names"]) == {INTUIT_A, INTUIT_B})
        assert pair["reason"] == "shared payroll"
        assert pair["names"][0] == INTUIT_B
        assert not any(OTHER in item["names"] and INTUIT_A in item["names"] for item in ideas)

        merge_vendors(conn, [INTUIT_A, INTUIT_B], "Intuit Payroll", actor="test")
        ideas = suggest_duplicates(spending, load_aliases(conn))
        assert not any(set(item["names"]) == {INTUIT_A, INTUIT_B} for item in ideas)

        merge_vendors(conn, ["ALPHA WIDGET SUPPLY", "BETA WIDGET SUPPLY"], "Widget Supply", actor="test")
        merge_vendors(conn, ["WIDGET DEPOT NORTH", "WIDGET DEPOT SOUTH"], "Widget Depot", actor="test")
        both = merge_vendors(conn, ["Widget Supply", "Widget Depot"], "Widget Co", actor="test")
        keys = {row["alias_key"]: row["canonical_name"] for row in both["aliases"]}
        assert keys["ALPHA WIDGET SUPPLY"] == "Widget Co"
        assert keys["WIDGET DEPOT SOUTH"] == "Widget Co"
        assert "Widget Co" not in keys


def _assert_canonical_pairs(ideas, aliases):
    merged = {key for key, target in aliases.items() if target != key}
    seen = set()
    for item in ideas:
        left, right = item["names"]
        assert left not in merged
        assert right not in merged
        assert left != right
        assert canonical_name(left, aliases) == left
        assert canonical_name(right, aliases) == right
        pair = frozenset(item["names"])
        assert pair not in seen
        seen.add(pair)


def test_suggestions_use_canonical_names_only(env):
    """Merged spellings roll up. A raw alias is never one side of a pair."""
    payroll = [f"INTUIT 83241316 DES:PAYROLL {index:02d}" for index in range(18)]
    decoy = "83241316 HOLDINGS"
    raw_prefix = "ACMECORP EXTRA DESCRIPTOR"
    aliases = {name: "Intuit Payroll" for name in payroll}
    aliases["Intuit Mid"] = "Intuit Payroll"
    aliases["INTUIT RAW CHAIN"] = "Intuit Mid"
    aliases[raw_prefix] = "Acme"
    spending = {name: 100 for name in payroll}
    spending["INTUIT RAW CHAIN"] = 50
    spending["Intuit Mid"] = 25
    spending["Intuit"] = 5000
    spending[decoy] = 700
    spending[raw_prefix] = 300
    spending["ACMECORP"] = 300
    # The canonical name can also appear as its own spending key.
    spending["Intuit Payroll"] = 40

    ideas = suggest_duplicates(spending, aliases)
    _assert_canonical_pairs(ideas, aliases)
    sides = {name for item in ideas for name in item["names"]}
    assert sides.isdisjoint(payroll)
    assert "Intuit Mid" not in sides
    assert "INTUIT RAW CHAIN" not in sides
    assert raw_prefix not in sides
    assert decoy not in sides
    assert "Acme" not in sides
    pairs = [item for item in ideas if set(item["names"]) == {"Intuit", "Intuit Payroll"}]
    assert len(pairs) == 1
    # 18 * 100 + chain 50 + mid 25 + the canonical key's own 40.
    assert pairs[0]["spend_cents"] == 5000 + 1800 + 50 + 25 + 40
    assert pairs[0]["names"][0] == "Intuit"
    assert pairs[0]["reason"] == "shared intuit"
    assert len(ideas) == 1

    # Same helper the Vendors page calls, against rows in the temporary database.
    with connect() as conn:
        for index, name in enumerate(payroll):
            _insert(conn, f"p{index}", "2026-06-02", -100, name)
        _insert(conn, "plain", "2026-06-03", -5000, "Intuit")
        _insert(conn, "decoy", "2026-06-04", -700, decoy)
        merge_vendors(conn, payroll, "Intuit Payroll", actor="test")
        from hpbooks.analytics import vendor_suggestions
        from hpbooks.vendors import load_aliases

        stored = load_aliases(conn)
        page = vendor_suggestions(conn, "2026-06-01", "2026-06-30", "all")
    _assert_canonical_pairs(page, stored)
    assert [set(item["names"]) for item in page] == [{"Intuit", "Intuit Payroll"}]
    assert page[0]["names"][0] == "Intuit"
    assert page[0]["spend_cents"] == 5000 + 18 * 100
    assert page[0]["reason"] == "shared intuit"
    assert not any(name in stored for item in page for name in item["names"])


def test_api_suggestions_omit_merged_spellings(env):
    payroll = [f"INTUIT 83241316 DES:PAYROLL {index:02d}" for index in range(18)]
    decoy = "83241316 HOLDINGS"
    with connect() as conn:
        for index, name in enumerate(payroll):
            _insert(conn, f"p{index}", "2026-06-02", -100, name)
        _insert(conn, "plain", "2026-06-03", -5000, "Intuit")
        _insert(conn, "decoy", "2026-06-04", -700, decoy)
        merge_vendors(conn, payroll, "Intuit Payroll", actor="test")

    from hpbooks.web import create_app

    client = create_app().test_client()
    page = client.get("/api/vendors?start=2026-06-01&end=2026-06-30").get_json()
    suggestions = page["suggestions"]
    raw = set(payroll)
    for item in suggestions:
        left, right = item["names"]
        assert left not in raw and right not in raw
        assert left != right
    pairs = [item for item in suggestions if set(item["names"]) == {"Intuit", "Intuit Payroll"}]
    assert len(pairs) == 1
    assert pairs[0]["names"][0] == "Intuit"
    assert pairs[0]["spend_cents"] == 5000 + 18 * 100
    assert pairs[0]["reason"] == "shared intuit"
    assert not any(decoy in item["names"] for item in suggestions)
    merchants = {row["merchant"] for row in page["rows"]}
    assert "Intuit Payroll" in merchants
    assert raw.isdisjoint(merchants)
    intuit = next(row for row in page["rows"] if row["merchant"] == "Intuit Payroll")
    assert intuit["spend_cents"] == 18 * 100
    assert intuit["alias_count"] == 18


def test_cli_merge_unmerge_and_aliases(env, capsys):
    with connect() as conn:
        _book(conn)
    assert main(["vendors", "merge", INTUIT_A, INTUIT_B, "--into", "Intuit Payroll"]) == 0
    assert "Intuit Payroll" in capsys.readouterr().out
    assert main(["vendors", "aliases"]) == 0
    listed = capsys.readouterr().out
    assert INTUIT_A in listed
    assert "Intuit Payroll" in listed
    assert main(["vendors", "merge", INTUIT_A, INTUIT_B, "--into", "Intuit Payroll"]) == 0
    assert "0 updated" in capsys.readouterr().out
    assert main(["vendors", "unmerge", INTUIT_A]) == 0
    assert "unmerged" in capsys.readouterr().out
    assert main(["vendors", "unmerge", INTUIT_A]) == 1
    assert "no alias" in capsys.readouterr().err
    with connect() as conn:
        count = conn.execute("SELECT COUNT(*) FROM audit_log WHERE action = 'vendor_merge'").fetchone()[0]
        assert count == 1


def test_api_filters_exports_and_csrf(env):
    from hpbooks.analytics import build_dashboard_range
    from hpbooks.web import create_app

    with connect() as conn:
        _book(conn)
    client = create_app().test_client()
    token = client.get("/api/session").get_json()["csrf_token"]

    def post(url, payload, csrf=token):
        headers = {"X-CSRF-Token": csrf} if csrf is not None else {}
        return client.post(url, json=payload, headers=headers)

    assert post("/api/vendors/merge", {"names": [INTUIT_A, INTUIT_B], "into": "Intuit Payroll"}, csrf=None).status_code == 403
    assert post("/api/vendors/merge", {"names": [INTUIT_A], "into": "Intuit Payroll"}).status_code == 400
    saved = post("/api/vendors/merge", {"names": [INTUIT_A, INTUIT_B], "into": "Intuit Payroll"})
    assert saved.status_code == 200
    assert saved.get_json()["changed"] == 2
    again = post("/api/vendors/merge", {"names": [INTUIT_A, INTUIT_B], "into": "Intuit Payroll"})
    assert again.get_json()["changed"] == 0

    page = client.get("/api/vendors?start=2026-06-01&end=2026-06-30").get_json()
    assert page["total_spend_cents"] == 3500
    intuit = next(row for row in page["rows"] if row["merchant"] == "Intuit Payroll")
    assert intuit["spend_cents"] == 3000
    assert intuit["alias_count"] == 2
    assert not any(item["names"] == [INTUIT_B, INTUIT_A] or set(item["names"]) == {INTUIT_A, INTUIT_B} for item in page["suggestions"])

    txns = client.get("/api/transactions?vendor=Intuit%20Payroll&start=2026-06-01&end=2026-06-30").get_json()
    assert {row["id"] for row in txns["rows"]} == {"a", "b"}
    assert {row["merchant_name"] for row in txns["rows"]} == {INTUIT_A, INTUIT_B}

    search = client.get("/api/search?q=Intuit%20Payroll").get_json()
    assert search["vendors"][0]["name"] == "Intuit Payroll"
    assert search["vendors"][0]["count"] == 2

    dash = client.get("/api/dashboard?start=2026-06-01&end=2026-06-30").get_json()
    top = {row["merchant"]: row["spend_cents"] for row in dash["top_vendors"]}
    assert top["Intuit Payroll"] == 3000
    assert INTUIT_A not in top

    csv_page = client.get("/export/expenses-by-vendor.csv?start=2026-06-01&end=2026-06-30&business=all")
    assert csv_page.status_code == 200
    text = csv_page.get_data(as_text=True)
    assert "Intuit Payroll" in text
    assert INTUIT_A not in text
    assert INTUIT_B not in text
    assert OTHER in text

    xlsx = client.get("/export/expenses-by-vendor.xlsx?start=2026-06-01&end=2026-06-30&business=all")
    assert xlsx.status_code == 200
    book = load_workbook(io.BytesIO(xlsx.data))
    sheet = book.active
    cells = [str(cell.value) for row in sheet.iter_rows() for cell in row if cell.value is not None]
    assert "Intuit Payroll" in cells
    assert INTUIT_A not in cells
    assert INTUIT_B not in cells

    pdf = client.get("/export/expenses-by-vendor.pdf?start=2026-06-01&end=2026-06-30&business=all")
    assert pdf.status_code == 200
    assert pdf.data.startswith(b"%PDF")
    if b"Intuit Payroll" in pdf.data:
        assert INTUIT_A.encode() not in pdf.data

    txn_csv = client.get("/export/transactions.csv?vendor=Intuit%20Payroll&start=2026-06-01&end=2026-06-30")
    assert txn_csv.status_code == 200
    exported = txn_csv.get_data(as_text=True)
    assert INTUIT_A in exported
    assert INTUIT_B in exported
    assert OTHER not in exported

    renamed = post("/api/vendors/rename", {"canonical": "Intuit Payroll", "name": "Payroll Processor"})
    assert renamed.status_code == 200
    assert renamed.get_json()["changed"] >= 1
    unmerged = post("/api/vendors/unmerge", {"alias": INTUIT_A})
    assert unmerged.status_code == 200
    page = client.get("/api/vendors?start=2026-06-01&end=2026-06-30").get_json()
    merchants = {row["merchant"]: row["spend_cents"] for row in page["rows"]}
    assert merchants[INTUIT_A] == 1000
    assert merchants["Payroll Processor"] == 2000

    with connect() as conn:
        actions = [row["action"] for row in conn.execute("SELECT action, actor FROM audit_log")]
        actors = {row["actor"] for row in conn.execute("SELECT actor FROM audit_log WHERE action LIKE 'vendor_%'")}
        assert actors == {"web"}
        assert "vendor_merge" in actions
        assert "vendor_unmerge" in actions
        assert "vendor_rename" in actions
        # The second merge did not add another audit row.
        assert actions.count("vendor_merge") == 1

    # Range helper still groups the same way.
    with connect() as conn:
        dash_rows = build_dashboard_range(conn, "2026-06-01", "2026-06-30", "all")["top_vendors"]
        assert any(row["merchant"] == "Payroll Processor" and row["spend_cents"] == 2000 for row in dash_rows)


def test_reads_work_when_the_alias_table_is_missing(env):
    with connect() as conn:
        _book(conn)
        conn.execute("DROP TABLE vendor_aliases")
    with connect(readonly=True) as conn:
        rows = _names(conn)
        assert {row["merchant"] for row in rows} == {INTUIT_A, INTUIT_B, OTHER}
        found = global_search(conn, "INTUIT")
        assert any(row["name"] == INTUIT_A for row in found["vendors"])
        from hpbooks.vendors import list_aliases

        assert list_aliases(conn) == []
