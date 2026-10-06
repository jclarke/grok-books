"""hpbooks tests. Always a temporary database and key — never the real ledger."""

from __future__ import annotations

import json
import sqlite3
from pathlib import Path

import pytest

from hpbooks.cli import main
from hpbooks.db import HpbooksError, connect, init_db, load_key
from fake_accounts import CHARGE_CARD, BANK, REWARDS
from hpbooks.importer import import_paths
from hpbooks.reports import build_pnl, reconcile

KEY = "ab" * 32  # 64 hex chars, not the real key


@pytest.fixture()
def env(tmp_path, monkeypatch):
    monkeypatch.setenv("HPBOOKS_KEY", KEY)
    monkeypatch.setenv("HPBOOKS_DB", str(tmp_path / "books.db"))
    monkeypatch.delenv("HPBOOKS_KEY_FILE", raising=False)
    init_db()
    return tmp_path


def _txn(txn_id, account, date, amount, name, pending="false", merchant=""):
    return {
        "id": txn_id,
        "account_id": account,
        "date": date,
        "amount": amount,
        "direction": "in" if not str(amount).startswith("-") else "out",
        "currency": "USD",
        "name": name,
        "merchant_name": merchant,
        "description": "",
        "pending": pending,
        "category": "",
    }


def _write_feed(path: Path, account: str, txns: list[dict], date_from="2026-01-01", date_to="2026-12-31"):
    path.write_text(
        json.dumps(
            {
                "source": "finance-mcp",
                "account_id": account,
                "date_from": date_from,
                "date_to": date_to,
                "fetched_at": "2026-03-01T00:00:00Z",
                "transactions": txns,
            }
        ),
        encoding="utf-8",
    )


def _class(txn_id):
    with connect() as conn:
        row = conn.execute(
            "SELECT * FROM classifications WHERE txn_id = ?", (txn_id,)
        ).fetchone()
        return None if row is None else {key: row[key] for key in row.keys()}


def _txn_row(txn_id):
    with connect() as conn:
        row = conn.execute("SELECT * FROM transactions WHERE id = ?", (txn_id,)).fetchone()
        return {key: row[key] for key in row.keys()}


def test_import_idempotent(env, tmp_path):
    feed = tmp_path / "feed.json"
    _write_feed(
        feed,
        BANK,
        [
            _txn("t-1", BANK, "2026-02-01", "10.00", "EXAMPLE PAYMENTS FUNDS DISB"),
            _txn("t-2", BANK, "2026-02-02", "-3.50", "Monthly Maintenance Fee"),
        ],
    )
    assert main(["import", str(feed)]) == 0
    assert main(["import", str(feed)]) == 0
    with connect() as conn:
        count = conn.execute("SELECT count(*) FROM transactions").fetchone()[0]
        logs = conn.execute(
            "SELECT inserted, updated, unchanged, superseded FROM import_log ORDER BY id"
        ).fetchall()
    assert count == 2
    assert logs[0]["inserted"] == 2
    assert logs[1]["inserted"] == 0
    assert logs[1]["updated"] == 0
    assert logs[1]["unchanged"] == 2
    assert logs[1]["superseded"] == 0


def test_pending_supersede_carries_manual_classification(env, tmp_path):
    first = tmp_path / "first.json"
    second = tmp_path / "second.json"
    _write_feed(
        first,
        BANK,
        [_txn("pend-1", BANK, "2026-03-10", "-25.00", "SOME VENDOR HOLD", pending="true")],
        date_from="2026-03-01",
        date_to="2026-03-31",
    )
    assert main(["import", str(first)]) == 0
    assert main(["classify", "pend-1", "--tag", "general", "--category", "Office/Other", "--note", "keep me"]) == 0

    _write_feed(
        second,
        BANK,
        [_txn("posted-1", BANK, "2026-03-12", "-25.00", "SOME VENDOR", pending="false")],
        date_from="2026-03-01",
        date_to="2026-03-31",
    )
    assert main(["import", str(second)]) == 0
    pending = _txn_row("pend-1")
    posted = _class("posted-1")
    assert pending["status"] == "superseded"
    assert pending["superseded_by"] == "posted-1"
    assert posted["source"] == "manual"
    assert posted["business_tag"] == "general"
    assert posted["category"] == "Office/Other"
    assert posted["note"] == "keep me"

    assert main(["import", str(second)]) == 0
    assert _txn_row("pend-1")["status"] == "superseded"
    with connect() as conn:
        last = conn.execute(
            "SELECT inserted, unchanged, superseded FROM import_log ORDER BY id DESC LIMIT 1"
        ).fetchone()
    assert last["inserted"] == 0
    assert last["superseded"] == 0
    assert last["unchanged"] == 1


def test_manual_classification_survives_reclassify_and_import(env, tmp_path):
    feed = tmp_path / "feed.json"
    _write_feed(
        feed,
        BANK,
        [_txn("man-1", BANK, "2026-04-01", "-80.00", "EXAMPLE PAYMENTS FUNDS DISB")],
    )
    assert main(["import", str(feed)]) == 0
    # The rule would call this a processor fee. Override it.
    assert main(["classify", "man", "--tag", "consulting", "--category", "Revenue - Consulting", "--note", "manual"]) == 0
    assert main(["reclassify", "--all"]) == 0
    assert main(["import", str(feed)]) == 0
    row = _class("man-1")
    assert row["source"] == "manual"
    assert row["business_tag"] == "consulting"
    assert row["category"] == "Revenue - Consulting"
    assert row["note"] == "manual"


def test_rule_creation_applies_to_other_rows(env, tmp_path):
    feed = tmp_path / "feed.json"
    _write_feed(
        feed,
        BANK,
        [
            _txn("w-1", BANK, "2026-05-01", "-12.00", "ACME WIDGETS 100"),
            _txn("w-2", BANK, "2026-05-02", "-12.00", "ACME WIDGETS 200"),
            _txn("w-3", BANK, "2026-05-03", "-15.00", "ACME WIDGETS 300"),
        ],
    )
    assert main(["import", str(feed)]) == 0
    assert (
        main(
            [
                "classify",
                "w-1",
                "--tag",
                "general",
                "--category",
                "Office/Other",
                "--note",
                "widgets",
                "--rule",
            ]
        )
        == 0
    )
    assert _class("w-1")["source"] == "manual"
    for txn_id in ("w-2", "w-3"):
        row = _class(txn_id)
        assert row["source"] == "rule"
        assert row["business_tag"] == "general"
        assert row["category"] == "Office/Other"
    with connect() as conn:
        count = conn.execute("SELECT count(*) FROM transactions").fetchone()[0]
    assert count == 3


def test_rule_application_audits_writes_and_skips_manual_and_superseded(env, tmp_path):
    feed = tmp_path / "widgets.json"
    _write_feed(
        feed,
        BANK,
        [
            _txn("rule-src", BANK, "2026-05-01", "-12.00", "QAWIDGET MERCHANT"),
            _txn("rule-other", BANK, "2026-05-02", "-12.00", "QAWIDGET MERCHANT"),
            _txn("rule-manual", BANK, "2026-05-03", "-12.00", "QAWIDGET MERCHANT"),
            _txn("rule-old", BANK, "2026-05-04", "-12.00", "QAWIDGET MERCHANT"),
        ],
    )
    assert main(["import", str(feed)]) == 0
    assert main(["classify", "rule-manual", "--tag", "consulting", "--category", "Revenue - Consulting", "--note", "leave me"]) == 0
    with connect() as conn:
        conn.execute("UPDATE transactions SET status = 'superseded' WHERE id = 'rule-old'")
    before = _class("rule-old")
    assert (
        main(
            [
                "classify",
                "rule-src",
                "--tag",
                "general",
                "--category",
                "Office/Other",
                "--note",
                "widgets",
                "--rule",
            ]
        )
        == 0
    )
    other = _class("rule-other")
    assert other["source"] == "rule"
    assert other["business_tag"] == "general"
    assert other["category"] == "Office/Other"
    manual = _class("rule-manual")
    assert manual["source"] == "manual"
    assert manual["business_tag"] == "consulting"
    assert manual["category"] == "Revenue - Consulting"
    assert _class("rule-old")["business_tag"] == before["business_tag"]
    assert _class("rule-old")["category"] == before["category"]
    assert _class("rule-old")["source"] == before["source"]
    with connect() as conn:
        applied_rows = conn.execute(
            "SELECT txn_id, action, rule_id, actor, note FROM audit_log WHERE note = 'rule applied'"
        ).fetchall()
        rule_id = conn.execute("SELECT id FROM rules WHERE pattern LIKE '%QAWIDGET%'").fetchone()["id"]
    assert [row["txn_id"] for row in applied_rows] == ["rule-other"]
    assert applied_rows[0]["action"] == "classify"
    assert applied_rows[0]["rule_id"] == rule_id
    assert applied_rows[0]["actor"] == "cli"
    from hpbooks.classify import apply_rule

    with connect() as conn:
        again = apply_rule(conn, rule_id, skip_txn_id="rule-src", actor="cli")
        extra = conn.execute("SELECT count(*) FROM audit_log WHERE note = 'rule applied'").fetchone()[0]
    assert again == 1
    assert extra == 1


def test_pnl_reconciles(env, tmp_path):
    feed = tmp_path / "feed.json"
    txns = [
        _txn("rev-1", BANK, "2026-01-15", "100.00", "EXAMPLE PAYMENTS FUNDS DISB"),
        _txn("cogs-1", BANK, "2026-01-16", "-40.00", "EXAMPLE COLO"),
        _txn("draw-1", BANK, "2026-01-17", "-10.00", "Online Banking transfer to CHK 0000"),
        _txn("unk-1", BANK, "2026-02-01", "-5.00", "ZZZ MYSTERY MERCHANT"),
        _txn("xfer-1", REWARDS, "2026-01-20", "10.00", "Payment Thank You-Mobile"),
    ]
    # One file can carry both accounts; window covers both.
    feed.write_text(
        json.dumps(
            {
                "source": "finance-mcp",
                "account_id": BANK,
                "date_from": "2026-01-01",
                "date_to": "2026-12-31",
                "transactions": txns,
            }
        ),
        encoding="utf-8",
    )
    assert main(["import", str(feed)]) == 0
    with connect() as conn:
        result = reconcile(conn, 2026)
        report = build_pnl(conn, 2026, by="month", business="all")
    assert result.ok, result.problems
    assert result.pnl_all == result.derived
    assert result.pnl_by_business["branda"] == 10000 - 4000  # revenue - cogs, in cents
    assert report.net_income_cents == result.pnl_all
    # Jan has revenue, cogs, draw, transfer. Feb has the unknown.
    labels = [row.label for row in report.rows]
    assert "Revenue - Hosting" in labels
    assert "Net Income" in labels
    assert "Owner Draws (below the line)" in labels
    assert "Net after Owner Draws" in labels
    assert "Memo: transfers excluded (1 rows, net)" in labels
    by_label = {row.label: row.values[-1] for row in report.rows}
    assert by_label["Owner Draws (below the line)"] == 1000  # -10.00 shown as a positive take
    assert by_label["Net Income"] - by_label["Owner Draws (below the line)"] == by_label["Net after Owner Draws"]
    assert by_label["Net Income"] == result.pnl_all
    csv_text = __import__("hpbooks.reports", fromlist=["pnl_csv"]).pnl_csv(report)
    assert "Owner Draws (below the line)" in csv_text
    assert "10.00" in csv_text
    assert report.columns[0] == "Jan 2026"
    assert report.columns[-1] == "Total"


def test_storage_modes_are_private(env, tmp_path):
    db = tmp_path / "books.db"
    db.chmod(0o644)
    tmp_path.chmod(0o755)
    with connect() as conn:
        conn.execute("SELECT 1").fetchone()
    assert db.stat().st_mode & 0o777 == 0o600
    assert tmp_path.stat().st_mode & 0o777 == 0o700


def test_database_unreadable_without_key(env, tmp_path):
    db = tmp_path / "books.db"
    conn = sqlite3.connect(db)
    with pytest.raises(sqlite3.DatabaseError, match="file is not a database"):
        conn.execute("SELECT count(*) FROM sqlite_master").fetchall()


def test_key_file_permissions(tmp_path, monkeypatch):
    monkeypatch.delenv("HPBOOKS_KEY", raising=False)
    key_path = tmp_path / "key"
    key_path.write_text(KEY + "\n", encoding="utf-8")
    key_path.chmod(0o644)
    monkeypatch.setenv("HPBOOKS_KEY_FILE", str(key_path))
    with pytest.raises(HpbooksError, match="too open"):
        load_key()
    key_path.chmod(0o600)
    assert load_key() == KEY


def test_csv_page_list_and_unknown_account(env, tmp_path, capsys):
    page = tmp_path / "pages.json"
    csv_body = (
        "id,account_id,date,amount,direction,currency,name,merchant_name,description,pending,category\n"
        f"csv-1,{BANK},2026-06-01,4.55,in,USD,STRIPE TRANSFER ST-1,,,false,\n"
        "csv-x,not-a-business-account,2026-06-01,-1.00,out,USD,Nope,,,false,\n"
    )
    page.write_text(
        json.dumps(
            [
                {"format": "csv", "csv": csv_body, "row_count": 2, "next_cursor": None},
            ]
        ),
        encoding="utf-8",
    )
    assert main(["import", str(page)]) == 0
    err = capsys.readouterr().err
    assert "unknown account_id" in err
    row = _class("csv-1")
    assert row["business_tag"] == "branda"
    assert row["category"] == "Revenue - Hosting"
    with connect() as conn:
        skipped = conn.execute("SELECT skipped FROM import_log").fetchone()[0]
        count = conn.execute("SELECT count(*) FROM transactions").fetchone()[0]
    assert skipped == 1
    assert count == 1


def test_owner_draw_category_is_enforced(env, tmp_path):
    feed = tmp_path / "feed.json"
    _write_feed(feed, BANK, [_txn("od-1", BANK, "2026-07-01", "-8.00", "ZZZ PERSONAL THING")])
    assert main(["import", str(feed)]) == 0
    assert main(["classify", "od-1", "--tag", "owner_draw", "--category", "Payroll"]) == 0
    row = _class("od-1")
    assert row["business_tag"] == "owner_draw"
    assert row["category"] == "Owner Draw"


def test_web_default_month_and_classify_api(env, tmp_path):
    feed = tmp_path / "months.json"
    _write_feed(
        feed,
        BANK,
        [
            _txn("old-rev", BANK, "2026-01-04", "-8.00", "JANUARY ONLY MERCHANT"),
            _txn("new-rev", BANK, "2026-09-04", "-9.00", "SEPTEMBER ONLY MERCHANT"),
            _txn("w-a", BANK, "2026-09-05", "-11.00", "ACME WIDGETS 11"),
            _txn("w-b", BANK, "2026-09-06", "-12.00", "ACME WIDGETS 22"),
            _txn("w-c", BANK, "2026-09-07", "-13.00", "ACME WIDGETS 33"),
            _txn("rev-sep", BANK, "2026-09-08", "20.00", "EXAMPLE PAYMENTS FUNDS DISB"),
        ],
    )
    assert main(["import", str(feed)]) == 0
    from hpbooks.web import create_app

    client = create_app().test_client()
    session = client.get("/api/session").get_json()
    assert session["latest_month"] == "2026-09"
    token = session["csrf_token"]

    def names(url):
        return {row["name"] for row in client.get(url).get_json()["rows"]}

    latest = names("/api/transactions?start=2026-09-01&end=2026-09-30")
    assert "SEPTEMBER ONLY MERCHANT" in latest
    assert "JANUARY ONLY MERCHANT" not in latest

    older = names("/api/transactions?month=2026-01")
    assert "JANUARY ONLY MERCHANT" in older
    assert "SEPTEMBER ONLY MERCHANT" not in older

    review = names("/api/transactions?month=2026-09&needs_review=1")
    assert "SEPTEMBER ONLY MERCHANT" in review
    assert not any("EXAMPLE PAYMENTS" in name for name in review)

    saved = client.post(
        "/api/classify",
        json={
            "txn_id": "w-a",
            "tag": "general",
            "category": "Office/Other",
            "note": "no rule",
            "save_rule": False,
        },
        headers={"X-CSRF-Token": token},
    )
    body = saved.get_json()
    assert saved.status_code == 200 and body["ok"] is True
    assert body["row"]["source"] == "manual"
    assert body["applied_others"] == 0
    assert _class("w-a")["source"] == "manual"
    assert _class("w-b")["source"] != "manual"

    ruled = client.post(
        "/api/classify",
        json={
            "txn_id": "w-b",
            "tag": "branda",
            "category": "Software & Licenses",
            "note": "widget rule",
            "save_rule": True,
            "pattern": r"ACME WIDGETS",
        },
        headers={"X-CSRF-Token": token},
    )
    ruled_body = ruled.get_json()
    assert ruled.status_code == 200 and ruled_body["ok"] is True
    assert ruled_body["applied_others"] >= 1
    assert _class("w-b")["source"] == "manual"
    assert _class("w-c")["source"] == "rule"
    assert _class("w-c")["category"] == "Software & Licenses"
    # The earlier manual row is not overwritten by the new rule.
    assert _class("w-a")["source"] == "manual"
    assert _class("w-a")["category"] == "Office/Other"


def test_exports_and_web(env, tmp_path):
    feed = tmp_path / "feed.json"
    _write_feed(
        feed,
        CHARGE_CARD,
        [_txn("fee-1", CHARGE_CARD, "2026-08-01", "-175.00", "RENEWAL MEMBERSHIP FEE")],
    )
    assert main(["import", str(feed)]) == 0
    xlsx = tmp_path / "pnl.xlsx"
    pdf = tmp_path / "pnl.pdf"
    assert main(["pnl", "--year", "2026", "--by", "month", "--format", "xlsx", "--out", str(xlsx)]) == 0
    assert main(["pnl", "--year", "2026", "--by", "year", "--format", "pdf", "--out", str(pdf)]) == 0
    assert xlsx.read_bytes()[:2] == b"PK"
    assert pdf.read_bytes()[:4] == b"%PDF"
    from openpyxl import load_workbook

    book = load_workbook(xlsx)
    assert "P&L" in book.sheetnames
    assert "Transactions" in book.sheetnames
    assert book["P&L"].freeze_panes == "B2"

    from hpbooks.web import create_app

    client = create_app().test_client()
    assert client.get("/healthz").get_json()["ok"] is True
    home = client.get("/api/transactions?month=2026-08")
    assert home.status_code == 200
    assert any(row["name"] == "RENEWAL MEMBERSHIP FEE" for row in home.get_json()["rows"])
    token = client.get("/api/session").get_json()["csrf_token"]
    pnl = client.get("/api/pnl?start=2026-01-01&end=2026-12-31&by=month&business=all")
    assert pnl.status_code == 200
    assert pnl.get_json()["title"] == "Example Co — Profit & Loss 2026-01-01 to 2026-12-31 (business: all)"
    for page in ("/", "/transactions", "/reports/pnl", "/review", "/rules", "/settings"):
        response = client.get(page)
        assert response.status_code == 200, page
        assert b'<div id="root">' in response.data
    assert client.get("/review").status_code == 200
    assert client.get("/rules").status_code == 200
    saved = client.post(
        "/api/classify",
        json={
            "txn_id": "fee-1",
            "tag": "general",
            "category": "Office/Other",
            "note": "from web",
        },
        headers={"X-CSRF-Token": token},
    )
    body = saved.get_json()
    assert saved.status_code == 200
    assert body["ok"] is True
    assert body["row"]["source"] == "manual"
    assert body["row"]["category"] == "Office/Other"
    exported = client.get("/export/pnl.csv?year=2026&by=month&business=all")
    assert exported.status_code == 200
    assert b"Net Income" in exported.data
    again = client.get("/api/transactions?month=2026-08").get_json()["rows"]
    fee = next(row for row in again if row["id"] == "fee-1")
    assert fee["note"] == "from web"
    assert fee["source"] == "manual"


def test_liability_csv_keeps_category_and_supersedes_empty_account(env, tmp_path):
    from fake_accounts import SAMPLE_CARD

    seeded = tmp_path / "seeded.json"
    seeded.write_text(
        json.dumps(
            {
                "source": "finance-mcp",
                "account_id": REWARDS,
                "date_from": "2026-09-01",
                "date_to": "2026-09-30",
                "transactions": [
                    _txn("card-1", REWARDS, "2026-09-25", "-6.00", "PADDLE.NET* SAMPLE ANALYTICS"),
                    _txn("charge-pend", CHARGE_CARD, "2026-09-23", "-3.00", "OLD HOLD", pending="true"),
                ],
            }
        ),
        encoding="utf-8",
    )
    # Give the rewards card row a provider category the later file will omit.
    payload = json.loads(seeded.read_text())
    payload["transactions"][0]["category"] = '{"primary":"GENERAL_SERVICES"}'
    seeded.write_text(json.dumps(payload), encoding="utf-8")
    assert main(["import", str(seeded)]) == 0

    pull = tmp_path / "cards.json"
    pull.write_text(
        json.dumps(
            {
                "format": "csv",
                "csv": (
                    "id,liability_id,date,amount,direction,currency,name,merchant_name,description,pending\n"
                    f"card-1,{REWARDS},2026-09-25,-6.0000,out,USD,PADDLE.NET* SAMPLE ANALYTICS,,,false\n"
                ),
                "row_count": 1,
                "next_cursor": None,
                "_query": {
                    "tool": "finance_query_liability_transactions",
                    "account_ids": [REWARDS, CHARGE_CARD, SAMPLE_CARD],
                    "date_from": "2026-09-22",
                    "date_to": "2026-09-28",
                },
            }
        ),
        encoding="utf-8",
    )
    assert main(["import", str(pull)]) == 0
    card = _txn_row("card-1")
    assert card["provider_category"] == '{"primary":"GENERAL_SERVICES"}'
    assert card["account_id"] == REWARDS
    pending = _txn_row("charge-pend")
    assert pending["status"] == "superseded"
    assert pending["superseded_by"] is None
    with connect() as conn:
        log = conn.execute(
            "SELECT inserted, updated, unchanged, superseded, skipped FROM import_log ORDER BY id DESC LIMIT 1"
        ).fetchone()
    assert log["inserted"] == 0
    assert log["updated"] == 0
    assert log["unchanged"] == 1
    assert log["superseded"] == 1
    assert log["skipped"] == 0


def test_unmatched_pending_is_still_superseded(env, tmp_path):
    first = tmp_path / "a.json"
    second = tmp_path / "b.json"
    _write_feed(
        first,
        BANK,
        [_txn("old-pend", BANK, "2026-04-02", "-9.00", "HOLD ITEM", pending="true")],
        date_from="2026-04-01",
        date_to="2026-04-30",
    )
    _write_feed(
        second,
        BANK,
        [_txn("other", BANK, "2026-04-20", "-90.00", "UNRELATED POSTED")],
        date_from="2026-04-01",
        date_to="2026-04-30",
    )
    with connect() as conn:
        import_paths(conn, [str(first)])
        import_paths(conn, [str(second)])
    row = _txn_row("old-pend")
    assert row["status"] == "superseded"
    assert row["superseded_by"] is None
