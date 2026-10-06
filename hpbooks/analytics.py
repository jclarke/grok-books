"""Dashboard, register, vendor, cash-flow, and Schedule C figures.

Amounts stay in integer cents. Expense figures are positive costs. Net income
matches build_pnl for the same transactions and business filter.
"""

from __future__ import annotations

import csv
import io
from dataclasses import dataclass, field
from datetime import date, timedelta
from pathlib import Path
from statistics import median
from xml.sax.saxutils import escape

from hpbooks.config import get_config
from hpbooks.db import (
    COGS_CATEGORIES,
    OPEX_CATEGORIES,
    REVENUE_CATEGORIES,
    HpbooksError,
    format_money,
    short_account,
)
from hpbooks.reports import (
    STATEMENT_CATEGORIES,
    _chmod_file,
    _column_label,
    _dollars,
    _load_txns,
    _plain_dollars,
    _tag,
    build_pnl,
    month_end,
    parse_month,
    prev_month,
    query_accounts,
    require_date,
)
from hpbooks.scope import SEED_ORDER, accounts_in, scope_clause

BUSINESS_FILTERS = ("all",) + get_config().business_slugs
BUSINESS_ERROR = get_config().business_filter_error


def revenue_label(category: str) -> str:
    """'Revenue - Sales' -> 'Sales' for charts."""
    prefix = "Revenue - "
    return category[len(prefix):] if category.startswith(prefix) else category


def _business_for_category(category: str) -> str:
    """The business a category suggests: the one whose revenue category it is, else the overhead business."""
    cfg = get_config()
    for item in cfg.businesses:
        if item.revenue_category == category:
            return item.slug
    return cfg.default_business


def _operating_label() -> str:
    cfg = get_config()
    operating = cfg.account_for_role("operating")
    for acct in cfg.accounts:
        if acct.id == operating:
            return acct.label or acct.short_name or acct.name
    return "operating account"


CHART_START = "2026-01"
OUTLOOK_DAYS = 60
# Monthly cadence: median gap between consecutive charges, in days.
MONTHLY_GAP_MIN = 25
MONTHLY_GAP_MAX = 35


@dataclass
class TabularCell:
    text: str
    cents: int | None = None


@dataclass
class TabularReport:
    title: str
    generated: str
    columns: list[str]
    rows: list[list[TabularCell]]
    kinds: list[str] = field(default_factory=list)


def _text(value) -> TabularCell:
    return TabularCell("" if value is None else str(value))


def _money(cents: int) -> TabularCell:
    return TabularCell(format_money(int(cents)), cents=int(cents))


def _csv_text(value: str) -> str:
    """Stop spreadsheet apps from treating a text cell as a formula."""
    value = value or ""
    if value[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + value
    return value


def period_figures(txns: list[dict], business: str) -> dict:
    """P&L figures for an already-loaded set of transactions."""

    def included(txn: dict) -> bool:
        if _tag(txn) in ("transfer", "owner_draw"):
            return False
        if business != "all" and _tag(txn) != business:
            return False
        return True

    income = [txn for txn in txns if included(txn)]

    def category_sum(category: str) -> int:
        return sum(
            int(txn["amount_cents"])
            for txn in income
            if _tag(txn) != "needs_review" and (txn.get("category") or "") == category
        )

    revenue = [(category, category_sum(category)) for category in REVENUE_CATEGORIES]
    refunds = category_sum("Refunds")
    net_revenue = sum(cents for _, cents in revenue) + refunds
    cogs_lines = [(category, -category_sum(category)) for category in COGS_CATEGORIES]
    cogs = sum(cents for _, cents in cogs_lines)
    opex: list[tuple[str, int]] = []
    opex_total = 0
    for category in OPEX_CATEGORIES:
        shown = -category_sum(category)
        if shown:
            opex.append((category, shown))
        opex_total += shown
    expenses: list[tuple[str, int]] = [(category, cents) for category, cents in cogs_lines if cents]
    expenses.extend(opex)
    plug = sum(
        int(txn["amount_cents"])
        for txn in income
        if _tag(txn) == "needs_review"
        or (
            _tag(txn) != "needs_review"
            and (txn.get("category") or "") not in STATEMENT_CATEGORIES
        )
    )
    net_income = sum(int(txn["amount_cents"]) for txn in income)
    if business == "all":
        owner_raw = sum(int(txn["amount_cents"]) for txn in txns if _tag(txn) == "owner_draw")
    else:
        owner_raw = 0
    owner_draws = -owner_raw
    return {
        "net_revenue": net_revenue,
        "total_expenses": cogs + opex_total,
        "net_income": net_income,
        "owner_draws": owner_draws,
        "net_after_draws": net_income - owner_draws,
        "revenue": revenue,
        "refunds": refunds,
        "cogs": cogs,
        "opex": opex,
        "expenses": expenses,
        "plug": plug,
    }


def _in_month(txns: list[dict], month: str) -> list[dict]:
    return [txn for txn in txns if (txn.get("date") or "").startswith(month)]


def _change(current: int, previous: int) -> tuple[int, float | None]:
    delta = current - previous
    if previous == 0:
        return delta, (0.0 if current == 0 else None)
    return delta, delta / abs(previous) * 100.0


def trend_months(anchor: str) -> list[str]:
    """Up to twelve months ending at anchor, none earlier than January 2026."""
    anchor = parse_month(anchor)
    if anchor < CHART_START:
        return [CHART_START]
    months = []
    cursor = anchor
    for _ in range(12):
        months.append(cursor)
        if cursor == CHART_START:
            break
        cursor = prev_month(cursor)
    months.reverse()
    return months


def _month_label(month: str) -> str:
    return _column_label(month).split(" ")[0]


def _ordered_accounts(rows: list[dict]) -> list[dict]:
    # The five seed accounts keep their usual order; other accounts follow.
    return sorted(rows, key=lambda row: SEED_ORDER.get(row["id"], 99))


def build_dashboard(conn, month: str, business: str) -> dict:
    month = parse_month(month)
    if business not in BUSINESS_FILTERS:
        raise HpbooksError(BUSINESS_ERROR)
    months = trend_months(month)
    previous = prev_month(month)
    load_from = min(months[0], previous) + "-01"
    txns = _load_txns(conn, load_from, month_end(month))
    current = period_figures(_in_month(txns, month), business)
    prior = period_figures(_in_month(txns, previous), business)
    labels = (
        ("net_revenue", "Net revenue"),
        ("total_expenses", "Total expenses"),
        ("net_income", "Net income"),
        ("owner_draws", "Owner draws"),
        ("net_after_draws", "Net after draws"),
    )
    kpis = []
    for key, label in labels:
        delta, pct = _change(current[key], prior[key])
        kpis.append(
            {
                "key": key,
                "label": label,
                "cents": current[key],
                "delta": delta,
                "pct": pct,
            }
        )
    points = []
    for item in months:
        figures = period_figures(_in_month(txns, item), business)
        points.append(
            {
                "month": item,
                "label": _month_label(item),
                "revenue": figures["net_revenue"],
                "expenses": figures["total_expenses"],
                "net": figures["net_income"],
            }
        )
    needs_review = review_count(conn)
    from hpbooks.balances import account_snapshots

    accounts = _ordered_accounts(query_accounts(conn))
    snaps = {snap["id"]: snap for snap in account_snapshots(conn)}
    split = [(revenue_label(category), cents) for category, cents in current["revenue"]]
    cash = [
        (short_account(row["name"]), int(snaps[row["id"]]["display_cents"]))
        for row in accounts
        if row["id"] in snaps
    ]
    return {
        "month": month,
        "business": business,
        "kpis": kpis,
        "plug": current["plug"],
        "points": points,
        "expenses": current["expenses"],
        "revenue_split": split,
        "cash": cash,
        "accounts": accounts,
        "review_count": needs_review,
        "prior_month": previous,
    }


def account_register(
    conn,
    account_id: str,
    *,
    search: str | None = None,
    date_from: str | None = None,
    date_to: str | None = None,
    limit: int = 100,
    offset: int = 0,
    status: str = "active",
) -> tuple[list[dict], int, int]:
    """Newest first. Running balance is the active-register balance after that row.

    Returns rows, filtered count, and the active ending balance.
    """
    if status not in ("active", "superseded", "all"):
        raise HpbooksError("status must be active, superseded, or all")
    from hpbooks.balances import running_balances, snapshot_for

    active = conn.execute(
        """
        SELECT id, date, amount_cents, pending
        FROM transactions
        WHERE account_id = ? AND status = 'active'
        ORDER BY date, id
        """,
        (account_id,),
    ).fetchall()
    snap = snapshot_for(conn, account_id)
    active_rows = [{key: row[key] for key in row.keys()} for row in active]
    running = running_balances(snap, active_rows)
    balance = int(snap["activity_cents"])
    clauses = ["t.account_id = ?"]
    params: list = [account_id]
    if status != "all":
        clauses.append("t.status = ?")
        params.append(status)
    if date_from:
        clauses.append("t.date >= ?")
        params.append(date_from)
    if date_to:
        clauses.append("t.date <= ?")
        params.append(date_to)
    if search:
        from hpbooks.reports import _like

        like = _like(search)
        clauses.append(
            "(ifnull(t.name,'') LIKE ? ESCAPE '\\' OR ifnull(t.merchant_name,'') LIKE ? ESCAPE '\\' "
            "OR ifnull(t.description,'') LIKE ? ESCAPE '\\' OR t.id LIKE ? ESCAPE '\\' "
            "OR ifnull(c.category,'') LIKE ? ESCAPE '\\')"
        )
        params.extend([like, like, like, like, like])
    where = " AND ".join(clauses)
    total = int(
        conn.execute(
            f"""
            SELECT COUNT(*)
            FROM transactions t
            LEFT JOIN classifications c ON c.txn_id = t.id
            WHERE {where}
            """,
            params,
        ).fetchone()[0]
    )
    rows = conn.execute(
        f"""
        SELECT t.id, t.date, t.amount_cents, t.name, t.merchant_name, t.description, t.pending,
               t.status, a.name AS account_name,
               c.business_tag, c.category, c.source, c.confidence, c.note
        FROM transactions t
        JOIN accounts a ON a.id = t.account_id
        LEFT JOIN classifications c ON c.txn_id = t.id
        WHERE {where}
        ORDER BY t.date DESC, t.id DESC
        LIMIT ? OFFSET ?
        """,
        [*params, int(limit), int(offset)],
    ).fetchall()
    result = []
    for row in rows:
        item = {key: row[key] for key in row.keys()}
        cents = int(item["amount_cents"])
        item["payment_cents"] = -cents if cents < 0 else None
        item["deposit_cents"] = cents if cents > 0 else None
        item["running_cents"] = running.get(item["id"])
        item["in_balance"] = item["id"] in running
        item["description"] = item.get("name") or item.get("merchant_name") or item.get("description") or ""
        result.append(item)
    return result, total, balance


def review_count(conn, mode: str = "business") -> int:
    """Business needs_review rows; in personal mode, the personal review queue."""
    if mode == "personal":
        from hpbooks.personal.classify import personal_review_count

        return personal_review_count(conn)
    scope_sql, scope_params = scope_clause(conn, "business")
    return int(
        conn.execute(
            f"""
            SELECT COUNT(*)
            FROM transactions t
            LEFT JOIN classifications c ON c.txn_id = t.id
            WHERE t.status = 'active'
              AND ifnull(c.business_tag, 'needs_review') = 'needs_review'
              AND {scope_sql}
            """,
            scope_params,
        ).fetchone()[0]
    )


def _business_ok(txn: dict, business: str) -> bool:
    if business == "all":
        return True
    return _tag(txn) == business


def _merchant_key(txn: dict) -> str:
    from hpbooks.vendors import vendor_key

    return vendor_key(txn.get("merchant_name"), txn.get("name"))


def _vendor_buckets(conn, start: str, end: str, business: str) -> dict[str, dict]:
    """Raw vendor key -> spend (positive cents), count, last date, category counts.

    Transfers, owner draws, and inflows are not vendor spend. Aliases are applied
    later so a merge never rewrites these keys.
    """
    txns = _load_txns(conn, start, end)
    groups: dict[str, dict] = {}
    for txn in txns:
        if _tag(txn) in ("transfer", "owner_draw"):
            continue
        if not _business_ok(txn, business):
            continue
        cents = int(txn["amount_cents"])
        if cents >= 0:
            continue
        key = _merchant_key(txn)
        bucket = groups.setdefault(key, {"spend": 0, "count": 0, "last": "", "cats": {}})
        bucket["spend"] += -cents
        bucket["count"] += 1
        if txn["date"] >= bucket["last"]:
            bucket["last"] = txn["date"]
        category = txn.get("category") or "Uncategorized"
        bucket["cats"][category] = bucket["cats"].get(category, 0) + 1
    return groups


def vendor_suggestions(conn, start: str, end: str, business: str) -> list[dict]:
    """Likely duplicate canonical vendors in this range.

    Raw spellings are rolled up through vendor aliases first. A suggestion
    never names an alias that is already merged, and two spellings of the
    same canonical vendor are not a pair.
    """
    from hpbooks.vendors import load_aliases, suggest_duplicates

    start, end = require_date(start), require_date(end)
    buckets = _vendor_buckets(conn, start, end, business)
    spending = {name: bucket["spend"] for name, bucket in buckets.items()}
    return suggest_duplicates(spending, load_aliases(conn))


def vendor_report(conn, start: str, end: str, business: str, limit: int = 25) -> tuple[list[dict], TabularReport]:
    start, end = require_date(start), require_date(end)
    if start > end:
        raise HpbooksError("start must be on or before end")
    if limit < 1:
        raise HpbooksError("limit must be positive")
    from hpbooks.vendors import canonical_name, load_aliases

    groups = _vendor_buckets(conn, start, end, business)
    aliases = load_aliases(conn)
    merged: dict[str, dict] = {}
    spellings: dict[str, list[str]] = {}
    alias_lists: dict[str, list[str]] = {}
    for raw, bucket in groups.items():
        canon = canonical_name(raw, aliases)
        dest = merged.setdefault(canon, {"spend": 0, "count": 0, "last": "", "cats": {}})
        dest["spend"] += bucket["spend"]
        dest["count"] += bucket["count"]
        if bucket["last"] >= dest["last"]:
            dest["last"] = bucket["last"]
        for category, count in bucket["cats"].items():
            dest["cats"][category] = dest["cats"].get(category, 0) + count
        spellings.setdefault(canon, []).append(raw)
        if raw in aliases:
            alias_lists.setdefault(canon, []).append(raw)
    for canon, bucket in merged.items():
        bucket["spellings"] = sorted(spellings.get(canon, []), key=str.lower)
        bucket["aliases"] = sorted(alias_lists.get(canon, []), key=str.lower)
    ordered = sorted(merged.items(), key=lambda item: (-item[1]["spend"], item[0].lower()))
    rows = []
    for name, bucket in ordered:
        category = sorted(bucket["cats"].items(), key=lambda item: (-item[1], item[0]))[0][0]
        rows.append(
            {
                "merchant": name,
                "spend_cents": bucket["spend"],
                "count": bucket["count"],
                "category": category,
                "last_seen": bucket["last"],
                "spellings": bucket.get("spellings") or [name],
                "aliases": bucket.get("aliases") or [],
                "alias_count": len(bucket.get("spellings") or [name]),
            }
        )
    shown = rows[:limit]
    rest = rows[limit:]
    cells = [
        [
            _text(row["merchant"]),
            _money(row["spend_cents"]),
            _text(str(row["count"])),
            _text(row["category"]),
            _text(row["last_seen"]),
        ]
        for row in shown
    ]
    kinds = ["line"] * len(shown)
    if rest:
        cells.append(
            [
                _text(f"Other vendors ({len(rest)})"),
                _money(sum(row["spend_cents"] for row in rest)),
                _text(str(sum(row["count"] for row in rest))),
                _text(""),
                _text(""),
            ]
        )
        kinds.append("subtotal")
    if rows:
        cells.append(
            [
                _text("Total"),
                _money(sum(row["spend_cents"] for row in rows)),
                _text(str(sum(row["count"] for row in rows))),
                _text(""),
                _text(""),
            ]
        )
        kinds.append("total")
    report = TabularReport(
        title=f"Expenses by vendor {start} to {end} (business: {business})",
        generated=date.today().isoformat(),
        columns=["Merchant", "Spend", "Count", "Category", "Last seen"],
        rows=cells,
        kinds=kinds,
    )
    return shown, report


def cash_flow_report(conn, start: str, end: str, business: str) -> TabularReport:
    start, end = require_date(start), require_date(end)
    if start > end:
        raise HpbooksError("start must be on or before end")
    txns = _load_txns(conn, start, end)
    buckets: dict[tuple[str, str], list[int]] = {}
    for txn in txns:
        if _tag(txn) == "transfer":
            continue
        if not _business_ok(txn, business):
            continue
        cents = int(txn["amount_cents"])
        key = (txn["date"][:7], txn.get("account_name") or "")
        bucket = buckets.setdefault(key, [0, 0])
        if cents >= 0:
            bucket[0] += cents
        else:
            bucket[1] += -cents
    rows = []
    kinds = []
    total_in = 0
    total_out = 0
    for (month, account), (money_in, money_out) in sorted(buckets.items()):
        total_in += money_in
        total_out += money_out
        rows.append(
            [
                _text(_column_label(month)),
                _text(short_account(account)),
                _money(money_in),
                _money(money_out),
                _money(money_in - money_out),
            ]
        )
        kinds.append("line")
    rows.append([_text("Total"), _text(""), _money(total_in), _money(total_out), _money(total_in - total_out)])
    kinds.append("total")
    return TabularReport(
        title=f"Cash flow {start} to {end} (business: {business}), transfers excluded",
        generated=date.today().isoformat(),
        columns=["Month", "Account", "Money in", "Money out", "Net"],
        rows=rows,
        kinds=kinds,
    )


def owner_draws_report(conn, start: str, end: str) -> TabularReport:
    start, end = require_date(start), require_date(end)
    if start > end:
        raise HpbooksError("start must be on or before end")
    txns = [txn for txn in _load_txns(conn, start, end) if _tag(txn) == "owner_draw"]
    rows = []
    kinds = []
    total = 0
    for txn in txns:
        taken = -int(txn["amount_cents"])
        total += taken
        rows.append(
            [
                _text(txn["date"]),
                _text(short_account(txn.get("account_name") or "")),
                _text(txn.get("name") or ""),
                _money(taken),
                _text(txn.get("class_note") or ""),
            ]
        )
        kinds.append("line")
    rows.append([_text("Total"), _text(""), _text(""), _money(total), _text("")])
    kinds.append("total")
    return TabularReport(
        title=f"Owner draws {start} to {end}",
        generated=date.today().isoformat(),
        columns=["Date", "Account", "Description", "Draw", "Note"],
        rows=rows,
        kinds=kinds,
    )


def pnl_by_business(conn, start: str, end: str) -> TabularReport:
    start, end = require_date(start), require_date(end)
    sides = get_config().business_slugs
    reports = {
        side: build_pnl(conn, int(start[:4]), by="year", business=side, start=start, end=end)
        for side in sides
    }
    combined = build_pnl(conn, int(start[:4]), by="year", business="all", start=start, end=end)
    side_maps = {
        side: {row.label: row.values[-1] for row in report.rows} for side, report in reports.items()
    }
    rows = []
    kinds = []
    for row in combined.rows:
        if row.kind == "memo":
            continue
        values = [side_maps[side].get(row.label, 0) for side in sides]
        # Owner-draw lines are company-wide; the side reports omit them.
        if row.label.startswith("Owner Draws") or row.label.startswith("Net after"):
            values = [0 for _ in sides]
        rows.append([_text(row.label)] + [_money(value) for value in values] + [_money(row.values[-1])])
        kinds.append(row.kind)
    return TabularReport(
        title=f"Profit & Loss by business {start} to {end}",
        generated=date.today().isoformat(),
        columns=["Line", *(business.label for business in get_config().businesses), "Total"],
        rows=rows,
        kinds=kinds,
    )


# Operating expense categories with a Schedule C line of their own, in line order. Every
# other configured operating expense category is listed under line 27 ("Other: <name>").
_SCHEDULE_C_LINES = {
    "Advertising": ("8", "Advertising"),
    "Contractors": ("11", "Contract labor"),
    "Insurance": ("15", "Insurance (other than health)"),
    "Interest": ("16b", "Interest (other)"),
    "Professional Services": ("17", "Legal and professional services"),
    "Office/Other": ("18", "Office expense"),
    "Rent": ("20b", "Rent or lease (other business property)"),
    "Taxes & Licenses": ("23", "Taxes and licenses"),
    "Travel": ("24a", "Travel"),
    "Meals": ("24b", "Deductible meals"),
    "Utilities": ("25", "Utilities"),
    "Payroll": ("26", "Wages"),
}


def _schedule_c_lines() -> list[tuple[str, str, str]]:
    """(line, label, category) for each configured operating expense category."""
    opex = get_config().opex_categories
    lines = [(code, name, category) for category, (code, name) in _SCHEDULE_C_LINES.items() if category in opex]
    others = [category for category in opex if category not in _SCHEDULE_C_LINES]
    for index, category in enumerate(others):
        code = "27" + ("abcdefghijklmnopqrstuvwxyz"[index] if index < 26 else "")
        lines.append((code, f"Other: {category}", category))
    return lines


def _cogs_line_label() -> str:
    """'Cost of goods sold', naming the cost category when there is one other than that."""
    cogs = get_config().cogs_categories
    if len(cogs) == 1 and cogs[0].lower() != "cost of goods sold":
        return f"Cost of goods sold ({cogs[0].lower()})"
    return "Cost of goods sold"


def schedule_c_report(conn, start: str, end: str, business: str) -> TabularReport:
    """Schedule C line order for bookkeeping. Not a tax return."""
    start, end = require_date(start), require_date(end)
    figures = period_figures(_load_txns(conn, start, end), business)
    opex = {label: cents for label, cents in figures["opex"]}
    line1 = sum(cents for _, cents in figures["revenue"])
    line2 = -figures["refunds"]
    line3 = line1 - line2
    line4 = figures["cogs"]
    line5 = line3 - line4
    line6 = 0
    line7 = line5 + line6
    expense_lines = [(code, name, opex.get(category, 0)) for code, name, category in _schedule_c_lines()]
    line28 = sum(cents for _, _, cents in expense_lines)
    line29 = line7 - line28
    rows = [
        [_text("1"), _text("Gross receipts or sales"), _money(line1)],
        [_text("2"), _text("Returns and allowances"), _money(line2)],
        [_text("3"), _text("Net receipts"), _money(line3)],
        [_text("4"), _text(_cogs_line_label()), _money(line4)],
        [_text("5"), _text("Gross profit"), _money(line5)],
        [_text("6"), _text("Other income"), _money(line6)],
        [_text("7"), _text("Gross income"), _money(line7)],
    ]
    kinds = ["line", "line", "subtotal", "line", "subtotal", "line", "subtotal"]
    for code, name, cents in expense_lines:
        rows.append([_text(code), _text(name), _money(cents)])
        kinds.append("line")
    rows.append([_text("28"), _text("Total expenses"), _money(line28)])
    kinds.append("subtotal")
    rows.append([_text("31"), _text("Net profit (or loss)"), _money(line29)])
    kinds.append("total")
    rows.append(
        [
            _text(""),
            _text("Memo: uncategorized / needs_review (not in the lines above)"),
            _money(figures["plug"]),
        ]
    )
    kinds.append("memo")
    rows.append(
        [
            _text(""),
            _text("Memo: net profit plus uncategorized equals P&L net income"),
            _money(line29 + figures["plug"]),
        ]
    )
    kinds.append("memo")
    rows.append(
        [
            _text(""),
            _text("Memo: owner draws are omitted (not a Schedule C deduction)"),
            _money(figures["owner_draws"]),
        ]
    )
    kinds.append("memo")
    return TabularReport(
        title=f"Schedule C-style summary {start} to {end} (business: {business})",
        generated=date.today().isoformat(),
        columns=["Line", "Description", "Amount"],
        rows=rows,
        kinds=kinds,
    )


def _gaps(dates: list[date]) -> list[int]:
    return [(dates[index] - dates[index - 1]).days for index in range(1, len(dates))]


def recurring_items(conn, *, business: str = "all", account_id: str | None = None) -> list[dict]:
    """Merchants with at least three charges spaced about a month apart."""
    txns = _load_txns(conn, "1900-01-01", "2999-12-31")
    groups: dict[tuple[str, str, str], list[dict]] = {}
    for txn in txns:
        if _tag(txn) == "transfer":
            continue
        if account_id and txn["account_id"] != account_id:
            continue
        if not _business_ok(txn, business):
            continue
        cents = int(txn["amount_cents"])
        if cents == 0:
            continue
        direction = "out" if cents < 0 else "in"
        key = (_merchant_key(txn), txn["account_id"], direction)
        if key[0] == "(unknown)":
            continue
        groups.setdefault(key, []).append(txn)
    found = []
    for (merchant, acct, direction), rows in groups.items():
        by_date: dict[str, int] = {}
        for txn in rows:
            by_date[txn["date"]] = by_date.get(txn["date"], 0) + int(txn["amount_cents"])
        days = sorted(date.fromisoformat(day) for day in by_date)
        if len(days) < 3:
            continue
        gaps = _gaps(days)
        gap = int(round(float(median(gaps))))
        if not (MONTHLY_GAP_MIN <= gap <= MONTHLY_GAP_MAX):
            continue
        last_day = days[-1]
        last_cents = by_date[last_day.isoformat()]
        next_day = last_day + timedelta(days=gap)
        account_name = rows[-1].get("account_name") or ""
        found.append(
            {
                "merchant": merchant,
                "account_id": acct,
                "account_name": short_account(account_name),
                "direction": direction,
                "count": len(rows),
                "last_date": last_day.isoformat(),
                "last_cents": last_cents,
                "gap_days": gap,
                "next_date": next_day.isoformat(),
            }
        )
    found.sort(key=lambda item: (item["next_date"], item["merchant"].lower()))
    return found


def cash_outlook(conn, *, business: str = "all", today: date | None = None, reserve_cents: int = 0) -> dict:
    """Estimate of the operating account's balance over the next 60 days.

    The operating account is the config account with role "operating".
    Opening balance is every active transaction in it. Projected events are
    its monthly deposits and withdrawals whose next date is still current.
    Card charges are not subtracted here; they leave the account only when the
    card is paid. The result is an estimate. With no operating account the
    outlook is flat at zero.
    """
    from hpbooks.balances import snapshot_for

    today = today or date.today()
    operating = get_config().account_for_role("operating")
    known = operating is not None and conn.execute(
        "SELECT 1 FROM accounts WHERE id = ?", (operating,)
    ).fetchone() is not None
    if known:
        snap = snapshot_for(conn, operating)
        # An anchor replaces the imported sum. Without one, the outlook still
        # starts from every active transaction in the account.
        balance = int(snap["real_balance_cents"] if snap["anchored"] else snap["activity_cents"])
        items = recurring_items(conn, business=business, account_id=operating)
    else:
        balance, items = 0, []
    horizon = today + timedelta(days=OUTLOOK_DAYS)
    events: dict[str, list[dict]] = {}
    for item in items:
        gap = int(item["gap_days"])
        cursor = date.fromisoformat(item["next_date"])
        if cursor < today - timedelta(days=gap + 3):
            continue
        guard = 0
        while cursor < today and guard < 24:
            cursor = cursor + timedelta(days=gap)
            guard += 1
        while cursor <= horizon and guard < 48:
            events.setdefault(cursor.isoformat(), []).append(
                {
                    "merchant": item["merchant"],
                    "cents": int(item["last_cents"]),
                    "direction": item["direction"],
                }
            )
            cursor = cursor + timedelta(days=gap)
            guard += 1
    running = balance
    timeline = [{"date": today.isoformat(), "label": "Opening balance", "cents": 0, "balance": running}]
    for day in sorted(events):
        for event in events[day]:
            running += int(event["cents"])
            timeline.append(
                {
                    "date": day,
                    "label": event["merchant"],
                    "cents": int(event["cents"]),
                    "balance": running,
                }
            )
    if not events:
        timeline.append(
            {
                "date": horizon.isoformat(),
                "label": f"No recurring {_operating_label()} items in this window",
                "cents": 0,
                "balance": running,
            }
        )
    daily = []
    cursor_balance = balance
    day = today
    while day <= horizon:
        key = day.isoformat()
        if key in events:
            cursor_balance += sum(int(event["cents"]) for event in events[key])
        daily.append({"date": key, "balance": cursor_balance})
        day += timedelta(days=1)
    return {
        "opening_cents": balance,
        "ending_cents": running,
        "reserve_cents": int(reserve_cents),
        "after_reserve_cents": running - int(reserve_cents),
        "timeline": timeline,
        "daily": daily,
        "estimate": True,
        "through": horizon.isoformat(),
    }


def table_csv(report: TabularReport) -> str:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(report.columns)
    for row in report.rows:
        writer.writerow(
            [
                _plain_dollars(cell.cents) if cell.cents is not None else _csv_text(cell.text)
                for cell in row
            ]
        )
    return buffer.getvalue()


def write_table_xlsx(report: TabularReport, path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    on_disk = isinstance(path, (str, Path))
    if on_disk:
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Report"
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F2937")
    subtotal_font = Font(bold=True)
    subtotal_fill = PatternFill("solid", fgColor="F3F4F6")
    total_font = Font(bold=True)
    total_fill = PatternFill("solid", fgColor="E5E7EB")
    memo_font = Font(italic=True, color="4B5563")
    currency = '$#,##0.00;($#,##0.00)'
    for col, header in enumerate(report.columns, start=1):
        cell = sheet.cell(1, col, header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="right" if col > 1 else "left")
    for row_index, row in enumerate(report.rows, start=2):
        kind = report.kinds[row_index - 2] if row_index - 2 < len(report.kinds) else "line"
        for col, cell_value in enumerate(row, start=1):
            value = _dollars(cell_value.cents) if cell_value.cents is not None else cell_value.text
            cell = sheet.cell(row_index, col, value)
            if cell_value.cents is not None:
                cell.number_format = currency
                cell.alignment = Alignment(horizontal="right")
            if kind == "subtotal":
                cell.font = subtotal_font
                cell.fill = subtotal_fill
            elif kind == "total":
                cell.font = total_font
                cell.fill = total_fill
            elif kind == "memo":
                cell.font = memo_font
    sheet.freeze_panes = "A2"
    last_col = get_column_letter(max(1, len(report.columns)))
    sheet.auto_filter.ref = f"A1:{last_col}{max(1, 1 + len(report.rows))}"
    for col in range(1, len(report.columns) + 1):
        sheet.column_dimensions[get_column_letter(col)].width = 18 if col > 1 else 42
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    workbook.properties.title = report.title
    workbook.properties.creator = "hpbooks"
    workbook.save(path)
    if on_disk:
        _chmod_file(path)


def write_table_pdf(report: TabularReport, path) -> None:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    on_disk = isinstance(path, (str, Path))
    if on_disk:
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        destination = str(path)
    else:
        destination = path
    doc = SimpleDocTemplate(
        destination,
        pagesize=landscape(letter),
        leftMargin=0.45 * inch,
        rightMargin=0.45 * inch,
        topMargin=0.45 * inch,
        bottomMargin=0.4 * inch,
        title=report.title,
    )
    styles = getSampleStyleSheet()
    title_style = styles["Title"]
    title_style.fontSize = 14
    title_style.leading = 17
    data = [report.columns]
    for row in report.rows:
        data.append([cell.text for cell in row])
    usable = landscape(letter)[0] - 0.9 * inch
    first = 2.4 * inch if len(report.columns) > 2 else usable * 0.4
    other = (usable - first) / max(1, len(report.columns) - 1)
    col_widths = [first] + [other] * (len(report.columns) - 1)
    table = Table(data, colWidths=col_widths, repeatRows=1)
    commands = [
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F2937")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("TEXTCOLOR", (0, 1), (-1, -1), colors.HexColor("#111827")),
        ("ALIGN", (1, 1), (-1, -1), "RIGHT"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("GRID", (0, 0), (-1, -1), 0.2, colors.HexColor("#D1D5DB")),
        ("LEFTPADDING", (0, 0), (-1, -1), 3),
        ("RIGHTPADDING", (0, 0), (-1, -1), 3),
        ("TOPPADDING", (0, 0), (-1, -1), 2),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
        ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
    ]
    for index, kind in enumerate(report.kinds, start=1):
        if kind == "subtotal":
            commands.append(("FONTNAME", (0, index), (-1, index), "Helvetica-Bold"))
            commands.append(("BACKGROUND", (0, index), (-1, index), colors.HexColor("#F3F4F6")))
        elif kind == "total":
            commands.append(("FONTNAME", (0, index), (-1, index), "Helvetica-Bold"))
            commands.append(("BACKGROUND", (0, index), (-1, index), colors.HexColor("#E5E7EB")))
        elif kind == "memo":
            commands.append(("TEXTCOLOR", (0, index), (-1, index), colors.HexColor("#4B5563")))
    table.setStyle(TableStyle(commands))
    story = [
        Paragraph(escape(report.title), title_style),
        Paragraph(escape(f"Generated {report.generated}"), styles["Normal"]),
        Spacer(1, 8),
        table,
    ]
    doc.build(story)
    if on_disk:
        _chmod_file(Path(path))


def transactions_csv(rows: list[dict]) -> str:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(
        ["date", "id", "account", "amount", "name", "merchant", "tag", "category", "source", "note", "status"]
    )
    for row in rows:
        writer.writerow(
            [
                row.get("date") or "",
                row.get("id") or "",
                _csv_text(row.get("account_name") or ""),
                _plain_dollars(int(row["amount_cents"])),
                _csv_text(row.get("name") or ""),
                _csv_text(row.get("merchant_name") or ""),
                row.get("business_tag") or "",
                _csv_text(row.get("category") or ""),
                row.get("source") or "",
                _csv_text(row.get("note") or ""),
                row.get("status") or "",
            ]
        )
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# JSON API figures for the single-page app. Everything stays in integer cents.


def compare_window(start: str, end: str) -> tuple[str, str]:
    """Prior window for a comparison.

    Whole calendar months compare to the same number of months just before.
    Other ranges use reports.prior_period.
    """
    from hpbooks.reports import prior_period

    start, end = require_date(start), require_date(end)
    start_d = date.fromisoformat(start)
    if start_d.day == 1 and end == month_end(end[:7]):
        months = (int(end[:4]) - start_d.year) * 12 + int(end[5:7]) - start_d.month + 1
        if 1 <= months < 12:
            first = start[:7]
            for _ in range(months):
                first = prev_month(first)
            return first + "-01", month_end(prev_month(start[:7]))
    return prior_period(start, end)


def _in_range(txns: list[dict], start: str, end: str) -> list[dict]:
    return [txn for txn in txns if start <= (txn.get("date") or "") <= end]


def upcoming_bills(conn, *, business: str = "all", today: date | None = None, days: int = 30) -> list[dict]:
    """Recurring outflows expected in the next `days` days, soonest first. An estimate."""
    today = today or date.today()
    horizon = today + timedelta(days=days)
    found = []
    for item in recurring_items(conn, business=business):
        if item["direction"] != "out":
            continue
        gap = int(item["gap_days"])
        due = date.fromisoformat(item["next_date"])
        if due < today - timedelta(days=gap + 3):
            continue  # stopped recurring
        guard = 0
        while due < today and guard < 24:
            due += timedelta(days=gap)
            guard += 1
        if due > horizon:
            continue
        found.append({**item, "due_date": due.isoformat(), "cents": int(item["last_cents"])})
    found.sort(key=lambda row: (row["due_date"], row["merchant"].lower()))
    return found


def build_dashboard_range(conn, start: str, end: str, business: str, *, today: date | None = None) -> dict:
    """Dashboard for any date range, compared with compare_window(start, end).

    Owner draws are left off the dashboard; they appear below the line in the P&L.
    """
    start, end = require_date(start), require_date(end)
    if start > end:
        raise HpbooksError("start must be on or before end")
    if business not in BUSINESS_FILTERS:
        raise HpbooksError(BUSINESS_ERROR)
    prior_start, prior_end = compare_window(start, end)
    months = trend_months(end[:7])
    load_from = min(months[0] + "-01", prior_start, start)
    # Chart points are whole calendar months. A range that ends mid-month still
    # draws a full bar for that month; the KPI figures stay inside start/end.
    load_to = month_end(end[:7])
    txns = _load_txns(conn, load_from, load_to)
    current = period_figures(_in_range(txns, start, end), business)
    prior = period_figures(_in_range(txns, prior_start, prior_end), business)

    def margin(fig: dict) -> float | None:
        if fig["net_revenue"] <= 0:
            return None
        return fig["net_income"] / fig["net_revenue"] * 100.0

    points = []
    for item in months:
        figures = period_figures(_in_month(txns, item), business)
        points.append(
            {
                "month": item,
                "label": _month_label(item),
                "revenue": figures["net_revenue"],
                "expenses": figures["total_expenses"],
                "net": figures["net_income"],
            }
        )
    kpis = []
    for key, label, spark in (
        ("net_revenue", "Net revenue", "revenue"),
        ("total_expenses", "Expenses", "expenses"),
        ("net_income", "Net income", "net"),
    ):
        delta, pct = _change(current[key], prior[key])
        kpis.append(
            {
                "key": key,
                "label": label,
                "cents": current[key],
                "prior_cents": prior[key],
                "delta": delta,
                "pct": pct,
                "spark": [point[spark] for point in points],
            }
        )
    current_margin, prior_margin = margin(current), margin(prior)
    from hpbooks.balances import account_snapshots

    accounts = _ordered_accounts(query_accounts(conn))
    snaps = {snap["id"]: snap for snap in account_snapshots(conn)}
    vendors, _table = vendor_report(conn, start, end, business, limit=6)
    return {
        "start": start,
        "end": end,
        "prior_start": prior_start,
        "prior_end": prior_end,
        "business": business,
        "kpis": kpis,
        "margin": {
            "pct": current_margin,
            "prior_pct": prior_margin,
            "delta": None if current_margin is None or prior_margin is None else current_margin - prior_margin,
        },
        "plug": current["plug"],
        "points": points,
        "expenses": [{"category": label, "cents": cents} for label, cents in current["expenses"]],
        "revenue_split": [
            {
                "label": revenue_label(item.revenue_category),
                "business": item.slug,
                "cents": dict(current["revenue"]).get(item.revenue_category, 0),
            }
            for item in get_config().businesses
            if item.revenue_category
        ],
        "refunds": current["refunds"],
        "cash": [
            {
                "id": row["id"],
                "name": row["name"],
                "short_name": short_account(row["name"]),
                "type": row["type"],
                "balance_cents": int(row["balance_cents"]),
                "display_cents": int(snaps[row["id"]]["display_cents"]) if row["id"] in snaps else int(row["balance_cents"]),
                "anchored": bool(snaps[row["id"]]["anchored"]) if row["id"] in snaps else False,
                "as_of_date": snaps[row["id"]]["as_of_date"] if row["id"] in snaps else None,
                "source": snaps[row["id"]]["source"] if row["id"] in snaps else None,
                "opening_cents": snaps[row["id"]]["opening_cents"] if row["id"] in snaps else None,
                "last_date": row["max_date"],
            }
            for row in accounts
        ],
        "top_vendors": vendors,
        "upcoming": upcoming_bills(conn, business=business, today=today),
        "review_count": review_count(conn),
    }


def accounts_range(conn, start: str, end: str) -> list[dict]:
    """The five accounts with balance, last activity, and in/out flow for the range."""
    from hpbooks.balances import attach_balances

    start, end = require_date(start), require_date(end)
    rows = _ordered_accounts(query_accounts(conn))
    flows = {
        row["account_id"]: (int(row["money_in"]), int(row["money_out"]))
        for row in conn.execute(
            """
            SELECT account_id,
                   COALESCE(SUM(CASE WHEN amount_cents > 0 THEN amount_cents ELSE 0 END), 0) AS money_in,
                   COALESCE(SUM(CASE WHEN amount_cents < 0 THEN -amount_cents ELSE 0 END), 0) AS money_out
            FROM transactions
            WHERE status = 'active' AND date >= ? AND date <= ? AND {scope_sql}
            GROUP BY account_id
            """.format(scope_sql=scope_clause(conn, "business", "account_id")[0]),
            (start, end, *scope_clause(conn, "business", "account_id")[1]),
        )
    }
    for row in rows:
        money_in, money_out = flows.get(row["id"], (0, 0))
        row["short_name"] = short_account(row["name"])
        row["in_cents"] = money_in
        row["out_cents"] = money_out
        row["net_cents"] = money_in - money_out
        row["balance_cents"] = int(row["balance_cents"])
    return attach_balances(conn, rows)


_PNL_GROUPS = {
    "Gross Revenue": REVENUE_CATEGORIES,
    "Net Revenue": REVENUE_CATEGORIES + ("Refunds",),
    "Total Operating Expenses": OPEX_CATEGORIES,
}


def pnl_line_transactions(
    conn, start: str, end: str, business: str, label: str, month: str | None = None
) -> list[dict]:
    """The transactions behind one P&L line, for drill-down. Same rules as build_pnl."""
    from hpbooks.db import CATEGORIES

    start, end = require_date(start), require_date(end)
    if month:
        month = parse_month(month)
        start = max(start, month + "-01")
        end = min(end, month_end(month))
    txns = _load_txns(conn, start, end)

    def in_income(txn: dict) -> bool:
        if _tag(txn) in ("transfer", "owner_draw"):
            return False
        return business == "all" or _tag(txn) == business

    income = [txn for txn in txns if in_income(txn)]

    def by_categories(categories) -> list[dict]:
        wanted = set(categories)
        return [
            txn
            for txn in income
            if _tag(txn) != "needs_review" and (txn.get("category") or "") in wanted
        ]

    if label in _PNL_GROUPS:
        rows = by_categories(_PNL_GROUPS[label])
    elif label == "Gross Profit":
        rows = by_categories(_PNL_GROUPS["Net Revenue"] + COGS_LINE)
    elif label in ("Net Income",):
        rows = income
    elif label.startswith("Uncategorized"):
        rows = [
            txn
            for txn in income
            if _tag(txn) == "needs_review" or (txn.get("category") or "") not in STATEMENT_CATEGORIES
        ]
    elif label.startswith("Owner Draws"):
        rows = [txn for txn in txns if _tag(txn) == "owner_draw"] if business == "all" else []
    elif label.startswith("Net after Owner Draws"):
        rows = income + [txn for txn in txns if _tag(txn) == "owner_draw"] if business == "all" else income
    elif label.startswith("Memo: transfers"):
        rows = [txn for txn in txns if _tag(txn) == "transfer"]
    elif label in CATEGORIES:
        rows = by_categories((label,))
    else:
        raise HpbooksError("unknown P&L line")
    rows = sorted(rows, key=lambda txn: (txn["date"], txn["id"]), reverse=True)
    return [
        {
            "id": txn["id"],
            "date": txn["date"],
            "account_id": txn["account_id"],
            "account_name": short_account(txn.get("account_name") or ""),
            "name": txn.get("name") or txn.get("merchant_name") or "",
            "amount_cents": int(txn["amount_cents"]),
            "business_tag": _tag(txn),
            "category": txn.get("category") or "",
        }
        for txn in rows
    ]


COGS_LINE = COGS_CATEGORIES


def review_suggestions(conn, rows: list[dict]) -> dict[str, dict]:
    """Suggested tag and category per review row, from how the same merchant was classified.

    Falls back to the category a needs_review rule already proposed.
    """
    from hpbooks.classify import default_pattern

    history: dict[str, dict[tuple[str, str], int]] = {}
    scope_sql, scope_params = scope_clause(conn, "business")
    for row in conn.execute(
        f"""
        SELECT t.name, t.merchant_name, c.business_tag, c.category
        FROM transactions t
        JOIN classifications c ON c.txn_id = t.id
        WHERE t.status = 'active' AND c.business_tag != 'needs_review' AND {scope_sql}
        """,
        scope_params,
    ):
        key = _merchant_key({"name": row["name"], "merchant_name": row["merchant_name"]}).lower()
        bucket = history.setdefault(key, {})
        pair = (row["business_tag"], row["category"])
        bucket[pair] = bucket.get(pair, 0) + 1
    out: dict[str, dict] = {}
    for row in rows:
        key = _merchant_key(row).lower()
        bucket = history.get(key)
        try:
            pattern = default_pattern(row)
        except HpbooksError:
            pattern = ""
        if bucket:
            (tag, category), count = sorted(bucket.items(), key=lambda item: (-item[1], item[0]))[0]
            total = sum(bucket.values())
            out[row["id"]] = {
                "tag": tag,
                "category": category,
                "reason": f"{count} of {total} earlier {_merchant_key(row)} rows",
                "confidence": round(count / total, 2),
                "pattern": pattern,
            }
            continue
        category = row.get("category") or ""
        if category and category not in ("Uncategorized", "Owner Draw", "Transfer"):
            tag = _business_for_category(category)
            out[row["id"]] = {
                "tag": tag,
                "category": category,
                "reason": "Category suggested by a review rule",
                "confidence": 0.4,
                "pattern": pattern,
            }
        else:
            out[row["id"]] = {
                "tag": "",
                "category": "",
                "reason": "",
                "confidence": 0,
                "pattern": pattern,
            }
    return out


def global_search(conn, query: str, limit: int = 8) -> dict:
    """Business transactions, vendors, and accounts whose text contains the query."""
    from hpbooks.reports import _like

    query = query.strip()
    if not query:
        return {"transactions": [], "vendors": [], "accounts": []}
    needle = query.lower()
    like = _like(query)
    scope_sql, scope_params = scope_clause(conn, "business")
    txns = conn.execute(
        f"""
        SELECT t.id, t.date, t.amount_cents, t.name, t.merchant_name, t.account_id, a.name AS account_name,
               ifnull(c.business_tag, 'needs_review') AS business_tag, c.category
        FROM transactions t
        JOIN accounts a ON a.id = t.account_id
        LEFT JOIN classifications c ON c.txn_id = t.id
        WHERE t.status = 'active'
          AND (ifnull(t.name,'') LIKE ? ESCAPE '\\' OR ifnull(t.merchant_name,'') LIKE ? ESCAPE '\\'
               OR ifnull(t.description,'') LIKE ? ESCAPE '\\' OR t.id LIKE ? ESCAPE '\\')
          AND {scope_sql}
        ORDER BY t.date DESC, t.id DESC
        LIMIT ?
        """,
        (like, like, like, like, *scope_params, int(limit)),
    ).fetchall()
    from hpbooks.vendors import VENDOR_SQL, combine_vendor_hits, load_aliases

    aliases = load_aliases(conn)
    vendors = conn.execute(
        f"""
        SELECT CASE WHEN trim(ifnull(t.merchant_name,'')) != '' THEN trim(t.merchant_name) ELSE trim(ifnull(t.name,'')) END AS vendor,
               COUNT(*) AS n,
               COALESCE(SUM(CASE WHEN t.amount_cents < 0 THEN -t.amount_cents ELSE 0 END), 0) AS spend,
               MAX(t.date) AS last_date
        FROM transactions t
        WHERE t.status = 'active'
          AND (ifnull(t.name,'') LIKE ? ESCAPE '\\' OR ifnull(t.merchant_name,'') LIKE ? ESCAPE '\\')
          AND {scope_sql}
        GROUP BY vendor
        ORDER BY n DESC, vendor
        LIMIT ?
        """,
        (like, like, *scope_params, max(int(limit) * 20, 50)),
    ).fetchall()
    hits = [
        {
            "name": row["vendor"],
            "count": int(row["n"]),
            "spend_cents": int(row["spend"]),
            "last_date": row["last_date"],
        }
        for row in vendors
        if row["vendor"]
    ]
    # A canonical name may not appear in the ledger text. Pull spellings whose
    # alias or display name contains the query so search still finds the group.
    needle_keys = sorted(
        {
            alias
            for alias, canon in aliases.items()
            if needle in alias.lower() or needle in canon.lower()
        }
    )
    seen = {row["name"] for row in hits}
    missing = [key for key in needle_keys if key not in seen]
    if missing:
        placeholders = ", ".join("?" for _ in missing)
        extra = conn.execute(
            f"""
            SELECT {VENDOR_SQL} AS vendor,
                   COUNT(*) AS n,
                   COALESCE(SUM(CASE WHEN t.amount_cents < 0 THEN -t.amount_cents ELSE 0 END), 0) AS spend,
                   MAX(t.date) AS last_date
            FROM transactions t
            WHERE t.status = 'active' AND {VENDOR_SQL} IN ({placeholders}) AND {scope_sql}
            GROUP BY vendor
            """,
            [*missing, *scope_params],
        ).fetchall()
        hits.extend(
            {
                "name": row["vendor"],
                "count": int(row["n"]),
                "spend_cents": int(row["spend"]),
                "last_date": row["last_date"],
            }
            for row in extra
            if row["vendor"]
        )
    accounts = [
        {"id": acct["id"], "name": acct["name"], "short_name": short_account(acct["name"])}
        for acct in accounts_in(conn, "business")
        if needle in acct["name"].lower() or needle in (acct["institution"] or "").lower()
    ]
    return {
        "transactions": [
            {
                "id": row["id"],
                "date": row["date"],
                "amount_cents": int(row["amount_cents"]),
                "name": row["name"] or row["merchant_name"] or "",
                "account_id": row["account_id"],
                "account_name": short_account(row["account_name"]),
                "business_tag": row["business_tag"],
                "category": row["category"] or "",
            }
            for row in txns
        ],
        "vendors": combine_vendor_hits(hits, aliases)[: int(limit)],
        "accounts": accounts,
    }


def tabular_json(report: TabularReport) -> dict:
    return {
        "title": report.title,
        "generated": report.generated,
        "columns": report.columns,
        "kinds": report.kinds,
        "rows": [[{"text": cell.text, "cents": cell.cents} for cell in row] for row in report.rows],
    }
