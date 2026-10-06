"""P&L, reconciliation, and list reports.

Net income is the sum of active non-transfer, non-owner-draw amounts in the
filter. Expense lines are shown as positive costs. The statement is built so
that figure ties out to the cent.
"""

from __future__ import annotations

import csv
import io
import os
import re
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import Decimal
from pathlib import Path
from xml.sax.saxutils import escape

from hpbooks.classify import as_dict, build_transfer_pairs
from hpbooks.config import get_config
from hpbooks.scope import has_scope, scope_clause
from hpbooks.db import short_id
from hpbooks.db import (
    BUSINESS_TAGS,
    CATEGORIES,
    COGS_CATEGORIES,
    OPEX_CATEGORIES,
    REVENUE_CATEGORIES,
    HpbooksError,
    account_name_map,
    format_money,
    short_account,
)

_DATE_RE = re.compile(r"^\d{4}-\d{2}-\d{2}$")
_MONTH_RE = re.compile(r"^\d{4}-\d{2}$")
TXN_SORTS = {
    "date": "t.date",
    "account": "a.name",
    "amount": "t.amount_cents",
    "name": "t.name",
    "tag": "ifnull(c.business_tag, '')",
    "category": "ifnull(c.category, '')",
    "status": "t.status",
}

STATEMENT_CATEGORIES = set(REVENUE_CATEGORIES) | set(COGS_CATEGORIES) | set(OPEX_CATEGORIES) | {"Refunds"}
BUSINESS_FILTERS = ("all",) + get_config().business_slugs
BUSINESS_ERROR = get_config().business_filter_error


@dataclass
class PnlRow:
    label: str
    kind: str  # line, subtotal, total, memo
    values: list[int]


@dataclass
class PnlReport:
    year: int
    by: str
    business: str
    ytd: bool
    title: str
    generated: str
    columns: list[str]
    rows: list[PnlRow]
    net_income_cents: int
    transfer_count: int
    transfer_cents: int
    owner_draw_cents: int
    detail: list[dict] = field(default_factory=list)
    prior_values: list[int] | None = None
    pct_values: list[float | None] | None = None
    prior_start: str = ""
    prior_end: str = ""


def require_date(value: str) -> str:
    """YYYY-MM-DD that is a real calendar date."""
    if not _DATE_RE.fullmatch(value or ""):
        raise HpbooksError("date must be YYYY-MM-DD")
    try:
        date.fromisoformat(value)
    except ValueError as exc:
        raise HpbooksError("date must be YYYY-MM-DD") from exc
    return value


def parse_month(value: str) -> str:
    """YYYY-MM with the month in 1..12."""
    if not _MONTH_RE.fullmatch(value or ""):
        raise HpbooksError("month must be YYYY-MM")
    year, month = int(value[:4]), int(value[5:7])
    if year < 1900 or year > 2200 or month < 1 or month > 12:
        raise HpbooksError("month must be YYYY-MM")
    return f"{year:04d}-{month:02d}"


def month_end(month: str) -> str:
    year, month_num = int(month[:4]), int(month[5:7])
    if month_num == 12:
        return f"{year:04d}-12-31"
    return (date(year, month_num + 1, 1) - timedelta(days=1)).isoformat()


def prev_month(month: str) -> str:
    year, month_num = int(month[:4]), int(month[5:7])
    if month_num == 1:
        return f"{year - 1:04d}-12"
    return f"{year:04d}-{month_num - 1:02d}"


def months_covering(start: str, end: str) -> list[str]:
    """Every calendar month touched by an inclusive date range, including empty ones."""
    start = require_date(start)
    end = require_date(end)
    if start > end:
        raise HpbooksError("start must be on or before end")
    year, month = int(start[:4]), int(start[5:7])
    last = end[:7]
    keys: list[str] = []
    while True:
        key = f"{year:04d}-{month:02d}"
        keys.append(key)
        if key == last:
            return keys
        month += 1
        if month == 13:
            year += 1
            month = 1
        if len(keys) > 600:
            raise HpbooksError("range is too long")


def prior_period(start: str, end: str) -> tuple[str, str]:
    """Window to compare against.

    A calendar year compares to the previous calendar year. A range that
    starts on January 1 compares to the same dates one year earlier. Every
    other range compares to the immediately preceding window of equal length.
    """
    start_d = date.fromisoformat(require_date(start))
    end_d = date.fromisoformat(require_date(end))
    if start_d.month == 1 and start_d.day == 1 and start_d.year == end_d.year:
        if end_d.month == 12 and end_d.day == 31:
            return f"{start_d.year - 1:04d}-01-01", f"{start_d.year - 1:04d}-12-31"
        try:
            prev_end = end_d.replace(year=end_d.year - 1)
        except ValueError:
            prev_end = date(end_d.year - 1, 2, 28)
        return f"{start_d.year - 1:04d}-01-01", prev_end.isoformat()
    span = (end_d - start_d).days + 1
    prior_end = start_d - timedelta(days=1)
    prior_start = prior_end - timedelta(days=span - 1)
    return prior_start.isoformat(), prior_end.isoformat()


def _year_bounds(year: int, ytd: bool, today: date) -> tuple[str, str]:
    start = f"{year:04d}-01-01"
    if ytd and year == today.year:
        end = today.isoformat()
    elif ytd and year > today.year:
        end = start
    else:
        end = f"{year:04d}-12-31"
    return start, end


def _load_txns(conn, start: str, end: str, mode: str = "business") -> list[dict]:
    scope_sql, scope_params = scope_clause(conn, mode)
    rows = conn.execute(
        f"""
        SELECT t.id, t.account_id, t.date, t.amount_cents, t.direction, t.currency,
               t.name, t.merchant_name, t.description, t.pending, t.status,
               a.name AS account_name,
               c.business_tag, c.category, c.source, c.confidence, c.note AS class_note
        FROM transactions t
        JOIN accounts a ON a.id = t.account_id
        LEFT JOIN classifications c ON c.txn_id = t.id
        WHERE t.status = 'active' AND t.date >= ? AND t.date <= ? AND {scope_sql}
        ORDER BY t.date, t.id
        """,
        (start, end, *scope_params),
    ).fetchall()
    return [as_dict(row) for row in rows]


def _tag(txn: dict) -> str:
    return txn.get("business_tag") or "needs_review"


def build_pnl(
    conn,
    year: int,
    by: str = "month",
    business: str = "all",
    ytd: bool = False,
    today: date | None = None,
    start: str | None = None,
    end: str | None = None,
) -> PnlReport:
    if business not in BUSINESS_FILTERS:
        raise HpbooksError(BUSINESS_ERROR)
    if by not in ("month", "year"):
        raise HpbooksError("by must be month or year")
    today = today or date.today()
    ranged = start is not None or end is not None
    if ranged:
        if not start or not end:
            raise HpbooksError("start and end are both required")
        start = require_date(start)
        end = require_date(end)
        if start > end:
            raise HpbooksError("start must be on or before end")
    else:
        start, end = _year_bounds(year, ytd, today)
    txns = _load_txns(conn, start, end)

    def in_business(txn: dict) -> bool:
        if business == "all":
            return True
        return _tag(txn) == business

    def in_income(txn: dict) -> bool:
        if _tag(txn) in ("transfer", "owner_draw"):
            return False
        return in_business(txn)

    ni_txns = [txn for txn in txns if in_income(txn)]
    if by == "year":
        keys = [str(year)]
    elif ranged:
        # Keep a column for every month in the requested window, even when it
        # has no transactions, so the statement covers the dates the user asked for.
        keys = months_covering(start, end)
    else:
        keys = sorted({txn["date"][:7] for txn in txns})
    col_keys = keys + ["Total"]

    def bucket(txn: dict) -> str:
        return str(year) if by == "year" else txn["date"][:7]

    def sum_amount(rows: list[dict]) -> list[int]:
        acc = {key: 0 for key in keys}
        for txn in rows:
            key = bucket(txn)
            if key in acc:
                acc[key] += int(txn["amount_cents"])
        total = sum(acc.values())
        return [acc[key] for key in keys] + [total]

    def add(left: list[int], right: list[int]) -> list[int]:
        return [a + b for a, b in zip(left, right)]

    def neg(values: list[int]) -> list[int]:
        return [-value for value in values]

    def category_rows(category: str) -> list[dict]:
        return [
            txn
            for txn in ni_txns
            if _tag(txn) != "needs_review" and txn.get("category") == category
        ]

    leaked = [
        txn
        for txn in ni_txns
        if _tag(txn) != "needs_review" and (txn.get("category") or "") not in STATEMENT_CATEGORIES
    ]
    review = [txn for txn in ni_txns if _tag(txn) == "needs_review"]

    revenue_rows = [PnlRow(category, "line", sum_amount(category_rows(category))) for category in REVENUE_CATEGORIES]
    gross = [0 for _ in col_keys]
    for row in revenue_rows:
        gross = add(gross, row.values)
    refunds = sum_amount(category_rows("Refunds"))
    net_rev = add(gross, refunds)
    cogs_rows = [PnlRow(category, "line", neg(sum_amount(category_rows(category)))) for category in COGS_CATEGORIES]
    cogs = [0 for _ in col_keys]
    for row in cogs_rows:
        cogs = add(cogs, row.values)
    gross_profit = [a - b for a, b in zip(net_rev, cogs)]

    opex_rows: list[PnlRow] = []
    opex_total = [0 for _ in col_keys]
    for category in OPEX_CATEGORIES:
        raw = sum_amount(category_rows(category))
        if all(value == 0 for value in raw):
            continue
        shown = neg(raw)
        opex_rows.append(PnlRow(category, "line", shown))
        opex_total = add(opex_total, shown)

    plug = sum_amount(review + leaked)
    ni = sum_amount(ni_txns)
    reconstructed = [gp - ox + pl for gp, ox, pl in zip(gross_profit, opex_total, plug)]
    if reconstructed != ni:
        raise HpbooksError("P&L lines do not add up to net income")

    columns = [_column_label(key) for key in col_keys]
    rows = [
        *revenue_rows,
        PnlRow("Gross Revenue", "subtotal", gross),
        PnlRow("Refunds", "line", refunds),
        PnlRow("Net Revenue", "subtotal", net_rev),
        *cogs_rows,
        PnlRow("Gross Profit", "subtotal", gross_profit),
        *opex_rows,
        PnlRow("Total Operating Expenses", "subtotal", opex_total),
    ]
    if business == "all" or any(value != 0 for value in plug):
        rows.append(PnlRow("Uncategorized / needs_review", "line", plug))
    rows.append(PnlRow("Net Income", "total", ni))

    owners = [txn for txn in txns if _tag(txn) == "owner_draw"]
    transfers = [txn for txn in txns if _tag(txn) == "transfer"]
    # Owner-draw amounts are outflows. Show them as a positive take, like expenses.
    owner_values = neg(sum_amount(owners))
    transfer_values = sum_amount(transfers)
    if business == "all":
        rows.append(PnlRow("Owner Draws (below the line)", "line", owner_values))
        rows.append(PnlRow("Net after Owner Draws", "total", [a - b for a, b in zip(ni, owner_values)]))
        rows.append(
            PnlRow(
                f"Memo: transfers excluded ({len(transfers)} rows, net)",
                "memo",
                transfer_values,
            )
        )

    detail_src = txns if business == "all" else ni_txns
    if ranged:
        title = f"{get_config().company_name} — Profit & Loss {start} to {end} (business: {business})"
    else:
        title = f"{get_config().company_name} — Profit & Loss {year} (business: {business})"
        if ytd:
            title += " YTD"
    return PnlReport(
        year=year,
        by=by,
        business=business,
        ytd=ytd,
        title=title,
        generated=today.isoformat(),
        columns=columns,
        rows=rows,
        net_income_cents=ni[-1],
        transfer_count=len(transfers),
        transfer_cents=transfer_values[-1],
        owner_draw_cents=owner_values[-1],
        detail=detail_src,
    )


def _column_label(key: str) -> str:
    if key == "Total":
        return "Total"
    if len(key) == 4 and key.isdigit():
        return key
    return datetime.strptime(key, "%Y-%m").strftime("%b %Y")


def render_table(headers: list[str], rows: list[list[str]], right_from: int = 1) -> str:
    widths = [len(header) for header in headers]
    for row in rows:
        for index, cell in enumerate(row):
            widths[index] = max(widths[index], len(cell))

    def format_row(row: list[str]) -> str:
        parts = []
        for index, cell in enumerate(row):
            if index < right_from:
                parts.append(cell.ljust(widths[index]))
            else:
                parts.append(cell.rjust(widths[index]))
        return "  ".join(parts)

    rule = "  ".join("-" * width for width in widths)
    lines = [format_row(headers), rule]
    lines.extend(format_row(row) for row in rows)
    return "\n".join(lines)


def render_pnl(report: PnlReport) -> str:
    headers = ["Line"] + report.columns
    body = []
    for row in report.rows:
        if row.kind == "memo" and (not body or not str(body[-1][0]).startswith("Owner")):
            body.append([""] * len(headers))
        body.append([row.label] + [format_money(value) for value in row.values])
    text = render_table(headers, body)
    return f"{report.title}\nGenerated {report.generated}\n\n{text}"


def _pct_text(value: float | None) -> str:
    if value is None:
        return ""
    return f"{value:.1f}"


def pnl_csv(report: PnlReport) -> str:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    headers = ["Line"] + report.columns
    compared = report.prior_values is not None
    if compared:
        headers += ["Prior", "% Change"]
    writer.writerow(headers)
    for index, row in enumerate(report.rows):
        values = [row.label] + [_plain_dollars(value) for value in row.values]
        if compared:
            values.append(_plain_dollars(report.prior_values[index]))
            pct = report.pct_values[index] if report.pct_values else None
            values.append(_pct_text(pct))
        writer.writerow(values)
    return buffer.getvalue()


def _plain_dollars(cents: int) -> str:
    sign = "-" if cents < 0 else ""
    cents = abs(int(cents))
    dollars, rem = divmod(cents, 100)
    return f"{sign}{dollars}.{rem:02d}"


def _dollars(cents: int) -> Decimal:
    return Decimal(int(cents)) / Decimal(100)


def _chmod_file(path: Path) -> None:
    os.chmod(path, 0o600)


def write_xlsx(report: PnlReport, path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    on_disk = isinstance(path, (str, Path))
    if on_disk:
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "P&L"
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F2937")
    subtotal_font = Font(bold=True)
    subtotal_fill = PatternFill("solid", fgColor="F3F4F6")
    total_font = Font(bold=True)
    total_fill = PatternFill("solid", fgColor="E5E7EB")
    memo_font = Font(italic=True, color="4B5563")
    currency = '$#,##0.00;($#,##0.00)'

    compared = report.prior_values is not None
    headers = ["Line"] + report.columns + (["Prior", "% Change"] if compared else [])
    for col, header in enumerate(headers, start=1):
        cell = sheet.cell(1, col, header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="right" if col > 1 else "left")
    for row_index, row in enumerate(report.rows, start=2):
        label_cell = sheet.cell(row_index, 1, row.label)
        if row.kind in ("subtotal", "total"):
            label_cell.font = total_font if row.kind == "total" else subtotal_font
        elif row.kind == "memo":
            label_cell.font = memo_font
        for col, value in enumerate(row.values, start=2):
            cell = sheet.cell(row_index, col, _dollars(value))
            cell.number_format = currency
            if row.kind == "subtotal":
                cell.font = subtotal_font
                cell.fill = subtotal_fill
                label_cell.fill = subtotal_fill
            elif row.kind == "total":
                cell.font = total_font
                cell.fill = total_fill
                label_cell.fill = total_fill
            elif row.kind == "memo":
                cell.font = memo_font
        if compared:
            prior_col = 2 + len(row.values)
            prior_cell = sheet.cell(row_index, prior_col, _dollars(report.prior_values[row_index - 2]))
            prior_cell.number_format = currency
            pct = report.pct_values[row_index - 2] if report.pct_values else None
            pct_cell = sheet.cell(row_index, prior_col + 1, None if pct is None else pct / 100.0)
            if pct is not None:
                pct_cell.number_format = "0.0%"
            if row.kind == "memo":
                prior_cell.font = memo_font
                pct_cell.font = memo_font
    sheet.freeze_panes = "B2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{1 + len(report.rows)}"
    sheet.column_dimensions["A"].width = 36
    for col in range(2, len(headers) + 1):
        sheet.column_dimensions[get_column_letter(col)].width = 16
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    workbook.properties.title = report.title
    workbook.properties.creator = "hpbooks"

    detail = workbook.create_sheet("Transactions")
    detail_headers = [
        "date",
        "id",
        "account",
        "amount",
        "name",
        "merchant_name",
        "business_tag",
        "category",
        "source",
        "confidence",
        "note",
        "pending",
    ]
    for col, header in enumerate(detail_headers, start=1):
        cell = detail.cell(1, col, header)
        cell.font = header_font
        cell.fill = header_fill
    for row_index, txn in enumerate(report.detail, start=2):
        values = [
            txn["date"],
            txn["id"],
            txn.get("account_name") or "",
            _dollars(int(txn["amount_cents"])),
            txn.get("name") or "",
            txn.get("merchant_name") or "",
            txn.get("business_tag") or "",
            txn.get("category") or "",
            txn.get("source") or "",
            txn.get("confidence"),
            txn.get("class_note") or "",
            int(txn.get("pending") or 0),
        ]
        for col, value in enumerate(values, start=1):
            cell = detail.cell(row_index, col, value)
            if col == 4:
                cell.number_format = currency
    detail.freeze_panes = "A2"
    detail.auto_filter.ref = f"A1:L{max(1, 1 + len(report.detail))}"
    for col, width in enumerate([12, 38, 28, 14, 48, 28, 18, 28, 12, 12, 40, 10], start=1):
        detail.column_dimensions[get_column_letter(col)].width = width
    workbook.save(path)
    if on_disk:
        _chmod_file(Path(path))


def write_pdf(report: PnlReport, path) -> None:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    on_disk = isinstance(path, (str, Path))
    if on_disk:
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        destination = str(path)
    else:
        destination = path
    doc = SimpleDocTemplate(
        destination,
        pagesize=landscape(letter),
        leftMargin=0.4 * inch,
        rightMargin=0.4 * inch,
        topMargin=0.45 * inch,
        bottomMargin=0.4 * inch,
        title=report.title,
    )
    styles = getSampleStyleSheet()
    title_style = styles["Title"]
    title_style.fontSize = 14
    title_style.leading = 17
    title_style.textColor = colors.HexColor("#111827")
    subtitle = styles["Normal"]
    subtitle.fontSize = 8
    subtitle.textColor = colors.HexColor("#4B5563")

    compared = report.prior_values is not None
    headers = ["Line"] + report.columns + (["Prior", "% Change"] if compared else [])
    data = [headers]
    for index, row in enumerate(report.rows):
        cells = [row.label] + [format_money(value) for value in row.values]
        if compared:
            cells.append(format_money(report.prior_values[index]))
            pct = report.pct_values[index] if report.pct_values else None
            cells.append("" if pct is None else f"{pct:.1f}%")
        data.append(cells)

    usable = landscape(letter)[0] - 0.8 * inch
    label_width = 2.15 * inch
    value_count = max(1, len(headers) - 1)
    other = (usable - label_width) / value_count
    col_widths = [label_width] + [other] * (len(headers) - 1)
    table = Table(data, colWidths=col_widths, repeatRows=1)
    commands = [
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 7),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F2937")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("TEXTCOLOR", (0, 1), (-1, -1), colors.HexColor("#111827")),
        ("ALIGN", (1, 0), (-1, -1), "RIGHT"),
        ("ALIGN", (0, 0), (0, -1), "LEFT"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("GRID", (0, 0), (-1, -1), 0.2, colors.HexColor("#D1D5DB")),
        ("LEFTPADDING", (0, 0), (-1, -1), 3),
        ("RIGHTPADDING", (0, 0), (-1, -1), 3),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
    ]
    for index, row in enumerate(report.rows, start=1):
        if row.kind == "subtotal":
            commands.append(("FONTNAME", (0, index), (-1, index), "Helvetica-Bold"))
            commands.append(("BACKGROUND", (0, index), (-1, index), colors.HexColor("#F3F4F6")))
        elif row.kind == "total":
            commands.append(("FONTNAME", (0, index), (-1, index), "Helvetica-Bold"))
            commands.append(("BACKGROUND", (0, index), (-1, index), colors.HexColor("#E5E7EB")))
        elif row.kind == "memo":
            commands.append(("TEXTCOLOR", (0, index), (-1, index), colors.HexColor("#4B5563")))
    table.setStyle(TableStyle(commands))
    story = [
        Paragraph(escape(report.title), title_style),
        Paragraph(escape(f"Generated {report.generated}"), subtitle),
        Spacer(1, 8),
        table,
    ]
    doc.build(story)
    if on_disk:
        _chmod_file(Path(path))


def write_pnl(report: PnlReport, fmt: str, path: Path | None) -> str:
    """Write or render a P&L. Returns text for table/csv, or the output path."""
    if fmt == "table":
        text = render_pnl(report)
        if path:
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_text(text + "\n", encoding="utf-8")
            _chmod_file(path)
        return text
    if fmt == "csv":
        text = pnl_csv(report)
        if path:
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_text(text, encoding="utf-8")
            _chmod_file(path)
        return text
    if fmt in ("xlsx", "pdf"):
        if path is None:
            raise HpbooksError(f"--out is required for {fmt}")
        if fmt == "xlsx":
            write_xlsx(report, path)
        else:
            write_pdf(report, path)
        return str(path)
    raise HpbooksError("format must be table, csv, xlsx, or pdf")


@dataclass
class ReconcileResult:
    year: int
    raw: int
    transfers: int
    transfer_count: int
    owner_draws: int
    owner_count: int
    derived: int
    pnl_all: int
    by_tag: dict[str, int]
    pnl_by_business: dict[str, int]
    review_line: int
    ok: bool
    problems: list[str]


def reconcile(conn, year: int) -> ReconcileResult:
    start, end = _year_bounds(year, False, date.today())
    txns = _load_txns(conn, start, end)
    raw = sum(int(txn["amount_cents"]) for txn in txns)
    transfers = [txn for txn in txns if _tag(txn) == "transfer"]
    owners = [txn for txn in txns if _tag(txn) == "owner_draw"]
    transfer_sum = sum(int(txn["amount_cents"]) for txn in transfers)
    owner_sum = sum(int(txn["amount_cents"]) for txn in owners)
    derived = raw - transfer_sum - owner_sum
    by_tag = {}
    for tag in (*get_config().business_slugs, "needs_review"):
        by_tag[tag] = sum(int(txn["amount_cents"]) for txn in txns if _tag(txn) == tag)
    pnl_all = build_pnl(conn, year, by="year", business="all")
    pnl_by_business = {}
    problems = []
    if pnl_all.net_income_cents != derived:
        problems.append(
            f"P&L net income {pnl_all.net_income_cents} != raw - transfers - owner draws {derived}"
        )
    for tag in get_config().business_slugs:
        report = build_pnl(conn, year, by="year", business=tag)
        pnl_by_business[tag] = report.net_income_cents
        if report.net_income_cents != by_tag[tag]:
            problems.append(f"{tag} P&L {report.net_income_cents} != tag sum {by_tag[tag]}")
    review_line = 0
    for row in pnl_all.rows:
        if row.label == "Uncategorized / needs_review":
            review_line = row.values[-1]
    tag_sum = sum(by_tag.values())
    if tag_sum != derived:
        problems.append(f"tag sums {tag_sum} != derived net income {derived}")
    return ReconcileResult(
        year=year,
        raw=raw,
        transfers=transfer_sum,
        transfer_count=len(transfers),
        owner_draws=owner_sum,
        owner_count=len(owners),
        derived=derived,
        pnl_all=pnl_all.net_income_cents,
        by_tag=by_tag,
        pnl_by_business=pnl_by_business,
        review_line=review_line,
        ok=not problems,
        problems=problems,
    )


def render_reconcile(result: ReconcileResult) -> str:
    lines = [
        f"Reconcile {result.year}",
        f"Raw sum of active txns:       {format_money(result.raw)}",
        f"Transfers ({result.transfer_count}):             {format_money(result.transfers)}",
        f"Owner draws ({result.owner_count}):            {format_money(result.owner_draws)}",
        f"Raw - transfers - draws:       {format_money(result.derived)}",
        f"P&L net income (all):          {format_money(result.pnl_all)}",
    ]
    for tag in get_config().business_slugs:
        lines.append(
            f"  {tag:<16} tag {format_money(result.by_tag[tag]):>14}   "
            f"P&L {format_money(result.pnl_by_business[tag]):>14}"
        )
    lines.append(f"  {'needs_review':<16} tag {format_money(result.by_tag['needs_review']):>14}")
    lines.append("OK" if result.ok else "FAIL")
    lines.extend(result.problems)
    return "\n".join(lines)


def distinct_months(conn, mode: str = "business") -> list[str]:
    scope_sql, scope_params = scope_clause(conn, mode, "account_id")
    rows = conn.execute(
        f"""
        SELECT DISTINCT substr(date, 1, 7) AS month
        FROM transactions
        WHERE status = 'active' AND {scope_sql}
        ORDER BY month
        """,
        scope_params,
    ).fetchall()
    return [row["month"] for row in rows]


def _like(value: str) -> str:
    escaped = value.replace("\\", "\\\\").replace("%", "\\%").replace("_", "\\_")
    return f"%{escaped}%"


def query_transactions(
    conn,
    *,
    month: str | None = None,
    account_id: str | None = None,
    tag: str | None = None,
    category: str | None = None,
    needs_review: bool = False,
    search: str | None = None,
    limit: int = 500,
    offset: int = 0,
    status: str = "active",
    min_cents: int | None = None,
    max_cents: int | None = None,
    sort: str = "date",
    direction: str = "desc",
    date_from: str | None = None,
    date_to: str | None = None,
    business: str | None = None,
    vendor: str | None = None,
    mode: str = "business",
) -> tuple[list[dict], int]:
    if status not in ("active", "superseded", "all"):
        raise HpbooksError("status must be active, superseded, or all")
    if sort not in TXN_SORTS:
        raise HpbooksError("unknown sort")
    if direction not in ("asc", "desc"):
        raise HpbooksError("direction must be asc or desc")
    if business not in (None, "", *BUSINESS_FILTERS):
        raise HpbooksError(BUSINESS_ERROR)
    clauses: list[str] = []
    params: list = []
    if status != "all":
        clauses.append("t.status = ?")
        params.append(status)
    scope_sql, scope_params = scope_clause(conn, mode)
    clauses.append(scope_sql)
    params.extend(scope_params)
    if month:
        clauses.append("substr(t.date, 1, 7) = ?")
        params.append(month)
    if date_from:
        clauses.append("t.date >= ?")
        params.append(date_from)
    if date_to:
        clauses.append("t.date <= ?")
        params.append(date_to)
    if account_id:
        clauses.append("t.account_id = ?")
        params.append(account_id)
    if tag:
        clauses.append("ifnull(c.business_tag, 'needs_review') = ?")
        params.append(tag)
    if business and business != "all":
        clauses.append("ifnull(c.business_tag, 'needs_review') = ?")
        params.append(business)
    if category:
        clauses.append("c.category = ?")
        params.append(category)
    if needs_review:
        clauses.append("ifnull(c.business_tag, 'needs_review') = 'needs_review'")
    if min_cents is not None:
        clauses.append("t.amount_cents >= ?")
        params.append(int(min_cents))
    if max_cents is not None:
        clauses.append("t.amount_cents <= ?")
        params.append(int(max_cents))
    if search:
        like = _like(search)
        clauses.append(
            "(ifnull(t.name,'') LIKE ? ESCAPE '\\' OR ifnull(t.merchant_name,'') LIKE ? ESCAPE '\\' "
            "OR ifnull(t.description,'') LIKE ? ESCAPE '\\' OR t.id LIKE ? ESCAPE '\\')"
        )
        params.extend([like, like, like, like])
    if vendor and vendor.strip():
        from hpbooks.vendors import vendor_filter_sql

        clause, vendor_params = vendor_filter_sql(conn, vendor)
        clauses.append(clause)
        params.extend(vendor_params)
    where = " AND ".join(clauses) if clauses else "1 = 1"
    total = conn.execute(
        f"""
        SELECT count(*)
        FROM transactions t
        LEFT JOIN classifications c ON c.txn_id = t.id
        WHERE {where}
        """,
        params,
    ).fetchone()[0]
    order_sql = TXN_SORTS[sort]
    order_dir = "ASC" if direction == "asc" else "DESC"
    rows = conn.execute(
        f"""
        SELECT t.id, t.date, t.amount_cents, t.name, t.merchant_name, t.description, t.pending,
               t.account_id, t.status, t.source AS txn_source, a.name AS account_name,
               c.business_tag, c.category, c.source, c.confidence, c.note
        FROM transactions t
        JOIN accounts a ON a.id = t.account_id
        LEFT JOIN classifications c ON c.txn_id = t.id
        WHERE {where}
        ORDER BY {order_sql} {order_dir}, t.id {order_dir}
        LIMIT ? OFFSET ?
        """,
        [*params, int(limit), int(offset)],
    ).fetchall()
    return [as_dict(row) for row in rows], int(total)


def query_review(conn, year: int | None = None, mode: str = "business") -> list[dict]:
    scope_sql, scope_params = scope_clause(conn, mode)
    clauses = [
        "t.status = 'active'",
        "ifnull(c.business_tag, 'needs_review') = 'needs_review'",
        scope_sql,
    ]
    params: list = list(scope_params)
    if year:
        clauses.append("substr(t.date, 1, 4) = ?")
        params.append(str(year))
    where = " AND ".join(clauses)
    rows = conn.execute(
        f"""
        SELECT t.id, t.date, t.amount_cents, t.name, t.merchant_name, t.description, t.pending,
               t.account_id, a.name AS account_name,
               c.business_tag, c.category, c.note, c.confidence, c.source
        FROM transactions t
        JOIN accounts a ON a.id = t.account_id
        LEFT JOIN classifications c ON c.txn_id = t.id
        WHERE {where}
        ORDER BY abs(t.amount_cents) DESC, t.date, t.id
        """,
        params,
    ).fetchall()
    return [as_dict(row) for row in rows]


def render_review(rows: list[dict], fmt: str = "table") -> str:
    if fmt == "csv":
        buffer = io.StringIO()
        writer = csv.writer(buffer)
        writer.writerow(["id", "date", "account", "amount", "name", "suggested_category", "note"])
        for row in rows:
            writer.writerow(
                [
                    row["id"],
                    row["date"],
                    row["account_name"],
                    _plain_dollars(int(row["amount_cents"])),
                    row.get("name") or "",
                    row.get("category") or "",
                    row.get("note") or "",
                ]
            )
        return buffer.getvalue()
    body = []
    for row in rows:
        body.append(
            [
                short_id(row["id"]),
                row["date"],
                short_account(row["account_name"]),
                format_money(int(row["amount_cents"])),
                (row.get("name") or "")[:64],
                row.get("category") or "",
                (row.get("note") or "")[:72],
            ]
        )
    headers = ["id", "date", "account", "amount", "name", "suggested", "note"]
    return render_table(headers, body, right_from=3)


def query_rules(conn) -> list[dict]:
    rows = conn.execute(
        "SELECT * FROM rules ORDER BY priority ASC, id ASC"
    ).fetchall()
    return [as_dict(row) for row in rows]


def query_accounts(conn, mode: str = "business") -> list[dict]:
    """Accounts in one mode with the imported-activity sum. Excluded accounts never appear."""
    where = "WHERE a.scope = ?" if has_scope(conn) else ("" if mode == "business" else "WHERE 0 = 1")
    params = [mode] if has_scope(conn) else []
    rows = conn.execute(
        f"""
        SELECT a.id, a.name, a.type, a.last4, a.institution,
               COALESCE(SUM(CASE WHEN t.status = 'active' THEN t.amount_cents ELSE 0 END), 0) AS balance_cents,
               COALESCE(SUM(CASE WHEN t.status = 'active' THEN 1 ELSE 0 END), 0) AS active_count,
               COALESCE(SUM(CASE WHEN t.status = 'superseded' THEN 1 ELSE 0 END), 0) AS superseded_count,
               MIN(CASE WHEN t.status = 'active' THEN t.date END) AS min_date,
               MAX(CASE WHEN t.status = 'active' THEN t.date END) AS max_date
        FROM accounts a
        LEFT JOIN transactions t ON t.account_id = a.id
        {where}
        GROUP BY a.id
        ORDER BY a.type, a.name
        """,
        params,
    ).fetchall()
    return [as_dict(row) for row in rows]


def render_accounts(rows: list[dict]) -> str:
    body = []
    for row in rows:
        body.append(
            [
                short_account(row["name"]),
                row["type"],
                row["last4"] or "",
                format_money(int(row["balance_cents"])),
                str(int(row["active_count"])),
                str(int(row["superseded_count"])),
                row["min_date"] or "",
                row["max_date"] or "",
            ]
        )
    return render_table(
        ["account", "type", "last4", "balance", "active", "superseded", "from", "to"],
        body,
        right_from=3,
    )


# Account boundary changes (register, scope move) show in each mode they touch.
_BOUNDARY_ACTIONS = ("account_scope", "account_register")


def query_audit(conn, limit: int = 200, mode: str = "all") -> list[dict]:
    """Recent audit rows. A mode keeps rows about that mode's transactions and accounts.

    Personal writes use actions prefixed personal_. A scope move or registration
    shows in the mode it came from and the mode it went to. Rows tied to no
    transaction or account (logins, settings, rules, WHMCS) belong to business.
    """
    where = "1 = 1"
    params: list = []
    if mode != "all" and has_scope(conn):
        if mode not in ("business", "personal"):
            raise HpbooksError("mode must be business, personal, or all")
        both = ", ".join("?" for _ in _BOUNDARY_ACTIONS)
        boundary = f"(action IN ({both}) AND (old_value = ? OR new_value = ?))"
        txn_in = "txn_id IN (SELECT t.id FROM transactions t JOIN accounts a ON a.id = t.account_id WHERE a.scope = ?)"
        acct_in = "field IN (SELECT id FROM accounts WHERE scope = ?)"
        if mode == "personal":
            where = f"({boundary} OR (action NOT IN ({both}) AND (action LIKE 'personal\\_%' ESCAPE '\\' OR {txn_in} OR {acct_in})))"
            params = [*_BOUNDARY_ACTIONS, "personal", "personal", *_BOUNDARY_ACTIONS, "personal", "personal"]
        else:
            where = (
                f"({boundary} OR (action NOT IN ({both}) AND action NOT LIKE 'personal\\_%' ESCAPE '\\' "
                f"AND (txn_id IS NULL OR {txn_in}) "
                f"AND (field IS NULL OR field NOT IN (SELECT id FROM accounts WHERE scope != ?))))"
            )
            params = [*_BOUNDARY_ACTIONS, "business", "business", *_BOUNDARY_ACTIONS, "business", "business"]
    rows = conn.execute(
        f"""
        SELECT id, ts, txn_id, rule_id, action, field, old_value, new_value, actor, note
        FROM audit_log
        WHERE {where}
        ORDER BY id DESC
        LIMIT ?
        """,
        (*params, int(limit)),
    ).fetchall()
    return [as_dict(row) for row in rows]


def render_audit(rows: list[dict]) -> str:
    body = []
    for row in rows:
        txn = short_id(row.get("txn_id"))
        body.append(
            [
                str(row["id"]),
                (row.get("ts") or "")[:19],
                row.get("action") or "",
                txn,
                "" if row.get("rule_id") is None else str(row["rule_id"]),
                row.get("actor") or "",
                (row.get("new_value") or row.get("note") or "")[:60],
            ]
        )
    return render_table(["id", "ts", "action", "txn", "rule", "actor", "detail"], body, right_from=0)


def transfer_rows(conn, mode: str = "business") -> list[dict]:
    scope_sql, scope_params = scope_clause(conn, mode, "account_id")
    txns = [
        as_dict(row)
        for row in conn.execute(
            f"SELECT * FROM transactions WHERE status = 'active' AND {scope_sql}", scope_params
        ).fetchall()
    ]
    classes = {
        row["txn_id"]: as_dict(row)
        for row in conn.execute("SELECT txn_id, business_tag, source, note FROM classifications")
    }
    names = account_name_map(conn)
    rows = []
    for item in build_transfer_pairs(txns):
        left = item["left"]
        right = item["right"]
        left_class = classes.get(left["id"]) if left else None
        right_class = classes.get(right["id"]) if right else None
        rows.append(
            {
                "kind": item["kind"],
                "status": item["status"],
                "left_date": left["date"] if left else "",
                "left_amount": int(left["amount_cents"]) if left else None,
                "left_account": short_account(names.get(left["account_id"], "")) if left else "",
                "left_name": (left.get("name") or "")[:48] if left else "",
                "left_id": left["id"] if left else "",
                "left_tag": (left_class or {}).get("business_tag") or "",
                "right_date": right["date"] if right else "",
                "right_amount": int(right["amount_cents"]) if right else None,
                "right_account": short_account(names.get(right["account_id"], "")) if right else "",
                "right_name": (right.get("name") or "")[:48] if right else "",
                "right_id": right["id"] if right else "",
                "right_tag": (right_class or {}).get("business_tag") or "",
                "note": (left_class or {}).get("note") or "",
            }
        )
    rows.sort(key=lambda row: (row["kind"], row["status"], row["left_date"] or row["right_date"], row["left_id"]))
    return rows


def render_transfers(rows: list[dict]) -> str:
    matched = sum(1 for row in rows if row["status"] == "matched")
    unmatched = len(rows) - matched
    body = []
    for row in rows:
        body.append(
            [
                row["kind"],
                row["status"],
                row["left_date"],
                "" if row["left_amount"] is None else format_money(row["left_amount"]),
                row["left_account"],
                row["left_tag"],
                row["right_date"],
                "" if row["right_amount"] is None else format_money(row["right_amount"]),
                row["right_account"],
                row["right_tag"],
                short_id(row["left_id"] or row["right_id"]),
            ]
        )
    table = render_table(
        ["kind", "status", "left_date", "left_amt", "left_acct", "left_tag", "right_date", "right_amt", "right_acct", "right_tag", "id"],
        body,
        right_from=3,
    )
    return f"Transfer pairs: {matched} matched, {unmatched} unmatched\n\n{table}"


def render_rules(rows: list[dict]) -> str:
    body = []
    for row in rows:
        account = row.get("account_id") or ""
        body.append(
            [
                str(row["id"]),
                str(row["priority"]),
                "yes" if row["active"] else "no",
                row.get("amount_sign") or "",
                account[:8],
                row["business_tag"],
                row["category"],
                f"{float(row['confidence']):.2f}",
                row["pattern"][:64],
            ]
        )
    return render_table(
        ["id", "pri", "on", "sign", "account", "tag", "category", "conf", "pattern"],
        body,
        right_from=0,
    )


def render_txns(rows: list[dict]) -> str:
    body = []
    for row in rows:
        body.append(
            [
                row["date"],
                short_account(row["account_name"]),
                format_money(int(row["amount_cents"])),
                (row.get("txn_source") or "")[:14],
                row.get("business_tag") or "",
                (row.get("category") or "")[:28],
                (row.get("name") or "")[:56],
                short_id(row["id"]),
            ]
        )
    return render_table(
        ["date", "account", "amount", "source", "tag", "category", "name", "id"],
        body,
        right_from=2,
    )


def resolve_account_arg(conn, text: str | None, scopes: tuple[str, ...] = ("business",)) -> str | None:
    """Account id from an id, last 4, or name fragment, among accounts in `scopes`."""
    if not text:
        return None
    scoped = has_scope(conn)
    allowed = [
        row["id"]
        for row in conn.execute(
            "SELECT id, scope FROM accounts" if scoped else "SELECT id, 'business' AS scope FROM accounts"
        )
        if row["scope"] in scopes
    ]
    if text in allowed:
        return text
    marks = ", ".join("?" for _ in allowed) or "NULL"
    rows = conn.execute(
        f"""
        SELECT id, name, last4 FROM accounts
        WHERE (last4 = ? OR id LIKE ? OR lower(name) LIKE ?) AND id IN ({marks})
        """,
        (text, text + "%", f"%{text.lower()}%", *allowed),
    ).fetchall()
    exact_last4 = [row for row in rows if row["last4"] == text]
    if len(exact_last4) == 1:
        return exact_last4[0]["id"]
    if len(rows) == 1:
        return rows[0]["id"]
    if not rows:
        raise HpbooksError(f"no account matches {text}")
    raise HpbooksError(f"{text} matches more than one account")


def allowed_tags() -> tuple[str, ...]:
    return BUSINESS_TAGS


def allowed_categories() -> tuple[str, ...]:
    return CATEGORIES


def _compare_label(label: str) -> str:
    """Stable key for prior-period matching.

    The transfer memo includes the current period's row count in its label
    ("Memo: transfers excluded (7 rows, net)"), so the text changes every
    period and would miss the prior line.
    """
    if label.startswith("Memo: transfers excluded"):
        return "Memo: transfers excluded"
    return label


def attach_prior(report: PnlReport, prior: PnlReport, start: str, end: str) -> PnlReport:
    """Add prior-period totals and percent change beside the report total column."""
    prior_map = {_compare_label(row.label): row.values[-1] for row in prior.rows}
    priors: list[int] = []
    pcts: list[float | None] = []
    for row in report.rows:
        current = int(row.values[-1])
        previous = int(prior_map.get(_compare_label(row.label), 0))
        priors.append(previous)
        if previous == 0:
            pcts.append(0.0 if current == 0 else None)
        else:
            pcts.append((current - previous) / abs(previous) * 100.0)
    report.prior_values = priors
    report.pct_values = pcts
    report.prior_start = start
    report.prior_end = end
    return report


def with_prior_period(
    conn,
    year: int,
    by: str = "month",
    business: str = "all",
    ytd: bool = False,
    today: date | None = None,
    start: str | None = None,
    end: str | None = None,
) -> PnlReport:
    """P&L for the requested window plus a prior-period comparison on the total."""
    today = today or date.today()
    report = build_pnl(
        conn,
        year,
        by=by,
        business=business,
        ytd=ytd,
        today=today,
        start=start,
        end=end,
    )
    if start is None and end is None:
        window_start, window_end = _year_bounds(year, ytd, today)
    else:
        window_start, window_end = start, end
    prior_start, prior_end = prior_period(window_start, window_end)
    prior = build_pnl(
        conn,
        year,
        by="year",
        business=business,
        ytd=False,
        today=today,
        start=prior_start,
        end=prior_end,
    )
    return attach_prior(report, prior, prior_start, prior_end)
